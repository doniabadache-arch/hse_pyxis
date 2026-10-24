"""
report_template.py — Modèle de rapport PROPRE À L'ENTREPRISE (HSE PYXIS)
=========================================================================
Permet à l'inspecteur d'importer le modèle de rapport de SA société ; l'IA rédige alors le
rapport selon ce modèle (mêmes titres, même ordre, mêmes colonnes de tableau).

Formats de modèle acceptés : .docx (recommandé), .pdf, .txt, .md

  * .docx : le fichier de sortie Word est construit SUR le modèle -> styles, polices,
            en-têtes / pieds de page, logo, marges de l'entreprise sont conservés.
            - Placez le paragraphe  {{REPORT_BODY}}  à l'endroit où le contenu doit être inséré.
              Ce qui PRÉCÈDE le repère (page de garde, sommaire...) reste intact ; ce qui le SUIT
              est l'exemple de structure : l'IA le suit puis il est remplacé par le rapport.
              Pour conserver une partie finale (signatures, annexes), ajoutez  {{END_REPORT_BODY}}
              après l'exemple : tout ce qui suit cette 2e balise est conservé.
            - Sans repère, tout le corps du modèle est remplacé par le rapport (en-têtes/pieds
              de page et styles conservés).
            - Variables remplacées partout (corps, en-têtes, pieds) :
              {{site}} {{inspection_id}} {{date}} {{inspector}} {{reviewer}} {{version}} {{ref}} {{title}}
               (anciens noms {{audit_id}} et {{auditor}} toujours acceptés)
  * .pdf / .txt / .md : servent de guide de structure pour l'IA ; la sortie Word utilise
            une mise en page standard.

Principe garde-fou (identique à report_render.py) : les LIGNES des tableaux d'écarts et tous
les chiffres viennent de la base (jamais du LLM). L'IA indique seulement quel tableau
"data_table" placer et quelles colonnes ; ce module remplit les valeurs. Si l'IA oublie le
tableau des écarts, il est ajouté automatiquement : les symptômes du Top 10 figurent donc
TOUJOURS tous dans le rapport.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

from report_render import _labels, _x, _as_list, _page, _REPORT_CSS, _logo_b64

TEMPLATES_DIR = Path(__file__).parent / "report_templates"
ALLOWED_EXT = (".docx", ".pdf", ".txt", ".md")
BODY_MARKER = "{{REPORT_BODY}}"
END_MARKER = "{{END_REPORT_BODY}}"
FALLBACK_STRUCTURE = (
    "(Aucune structure fournie après le repère : utilise la structure standard d'un rapport d'inspection HSE — "
    "1. Synthèse exécutive, 2. Analyse quantitative, 3. Risques prioritaires, 4. Analyse par domaine, "
    "5. Recommandations, 6. Conclusion et prochaines étapes.)")
MAX_OUTLINE_CHARS = 12000
MAX_TEMPLATE_BYTES = 15 * 1024 * 1024

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ======================================================================
# 1. STOCKAGE DES MODÈLES (réutilisables d'une inspection à l'autre)
# ======================================================================
def _safe_name(name: str) -> str:
    base = Path(str(name)).name
    base = re.sub(r"[^\w.\- ]+", "_", base, flags=re.U).strip(" .")
    return base or "modele"


def list_saved_templates() -> list[str]:
    if not TEMPLATES_DIR.exists():
        return []
    return sorted(p.name for p in TEMPLATES_DIR.iterdir()
                  if p.is_file() and p.suffix.lower() in ALLOWED_EXT)


def save_template(filename: str, data: bytes) -> str:
    name = _safe_name(filename)
    if Path(name).suffix.lower() not in ALLOWED_EXT:
        raise ValueError("Format non pris en charge : " + ", ".join(ALLOWED_EXT))
    if len(data) > MAX_TEMPLATE_BYTES:
        raise ValueError("Modèle trop volumineux (15 Mo max).")
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    (TEMPLATES_DIR / name).write_bytes(data)
    return name


def load_saved_template(name: str) -> bytes | None:
    path = TEMPLATES_DIR / _safe_name(name)
    return path.read_bytes() if path.is_file() else None


def delete_saved_template(name: str) -> None:
    path = TEMPLATES_DIR / _safe_name(name)
    if path.is_file():
        path.unlink()


# ======================================================================
# 2. EXTRACTION DE LA STRUCTURE DU MODÈLE (pour l'IA)
# ======================================================================
_HEAD_RE = re.compile(r"^\s*(?:heading|titre|عنوان)\s*(\d)\s*$", re.I)


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1256", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _outline_docx(data: bytes) -> dict:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    children = list(doc.element.body.iterchildren())
    has_marker = any(
        el.tag.rsplit("}", 1)[-1] == "p" and BODY_MARKER in Paragraph(el, doc).text for el in children)

    # Avec repère : seule la zone « exemple de structure » (entre {{REPORT_BODY}} et
    # {{END_REPORT_BODY}} / fin du document) est envoyée à l'IA ; le reste est conservé tel quel.
    state = 1 if not has_marker else 0   # 0 = avant, 1 = dans la zone, 2 = après
    lines, headings = [], []
    for el in children:
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            txt = p.text.strip()
            if has_marker and BODY_MARKER in txt:
                state = 1
                continue
            if has_marker and END_MARKER in txt:
                state = 2
                continue
            if state != 1 or not txt:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = _HEAD_RE.match(style)
            if m:
                lvl = max(1, min(3, int(m.group(1))))
                lines.append("#" * lvl + " " + txt)
                headings.append((lvl, txt))
            elif style.strip().lower() in ("title", "titre", "titre principal"):
                lines.append("# " + txt)
                headings.append((1, txt))
            else:
                lines.append(txt)
        elif tag == "tbl" and state == 1:
            tbl = Table(el, doc)
            rows = []
            for r in tbl.rows[:4]:
                rows.append(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells))
            if rows:
                lines.append("[TABLEAU] colonnes : " + rows[0])
                for ex in rows[1:]:
                    if ex.replace("|", "").strip():
                        lines.append("[TABLEAU] exemple de ligne : " + ex)
    outline = "\n".join(lines)[:MAX_OUTLINE_CHARS]
    if has_marker and not outline.strip():
        outline = FALLBACK_STRUCTURE
    return {"ext": ".docx", "outline": outline, "headings": headings, "has_marker": has_marker}


def _outline_pdf(data: bytes) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
    return {"ext": ".pdf", "outline": text.strip()[:MAX_OUTLINE_CHARS],
            "headings": _guess_headings(text), "has_marker": False}


def _guess_headings(text: str) -> list[tuple[int, str]]:
    out = []
    for ln in text.splitlines():
        s = ln.strip()
        if not s or len(s) > 90:
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            out.append((len(m.group(1)), m.group(2)))
        elif re.match(r"^\d+(\.\d+)*[.)]?\s+\S", s):
            out.append((min(3, s.split()[0].count(".") + 1), s))
    return out[:60]


def extract_outline(data: bytes, filename: str) -> dict:
    """-> {"ext", "outline" (texte pour l'IA), "headings" [(niveau, titre)], "has_marker"}.
    Lève ValueError si le fichier est illisible."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ValueError("Format non pris en charge : " + ", ".join(ALLOWED_EXT))
    try:
        if ext == ".docx":
            return _outline_docx(data)
        if ext == ".pdf":
            return _outline_pdf(data)
    except Exception as e:  # fichier corrompu / protégé
        raise ValueError(f"Impossible de lire le modèle ({e}).") from e
    text = _decode_text(data)
    return {"ext": ext, "outline": text.strip()[:MAX_OUTLINE_CHARS],
            "headings": _guess_headings(text), "has_marker": BODY_MARKER in text}


# ======================================================================
# 3. NORMALISATION DE LA SORTIE IA + GARANTIE « TOUS LES SYMPTÔMES »
# ======================================================================
def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def normalize_report(raw: dict, items: list, lang: str = "fr") -> dict:
    """Nettoie le JSON de l'IA et garantit la présence du tableau de TOUS les écarts."""
    L = _labels(lang)
    raw = raw if isinstance(raw, dict) else {}

    sections = []
    for s in _as_list(raw.get("sections")):
        if not isinstance(s, dict):
            continue
        heading = _txt(s.get("heading"))
        try:
            level = max(1, min(3, int(s.get("level", 1))))
        except (TypeError, ValueError):
            level = 1
        kind = _txt(s.get("data_table")).lower()
        if kind not in ("priority_risks", "compliance_summary", "domain_breakdown"):
            kind = ""
        tbl = s.get("table") if isinstance(s.get("table"), dict) else None
        sections.append({
            "heading": heading,
            "level": level,
            "paragraphs": [_txt(p) for p in _as_list(s.get("paragraphs")) if _txt(p)],
            "bullets": [_txt(b) for b in _as_list(s.get("bullets")) if _txt(b)],
            "data_table": kind,
            "columns": [_txt(c).lower() for c in _as_list(s.get("columns"))],
            "headers": [_txt(h) for h in _as_list(s.get("headers"))],
            "table": tbl,
        })

    notes = {}
    for n in _as_list(raw.get("risk_notes")):
        if isinstance(n, dict) and n.get("item_id") is not None:
            notes[str(n["item_id"])] = {
                "why_critical": _txt(n.get("why_critical")),
                "potential_impact": _txt(n.get("potential_impact")),
            }

    # Garantie : si l'IA n'a prévu aucun tableau des écarts, on l'ajoute (toutes les lignes viennent de la BD).
    if items and not any(s["data_table"] == "priority_risks" for s in sections):
        sections.append({"heading": L["s3"], "level": 1, "paragraphs": [], "bullets": [],
                         "data_table": "priority_risks", "columns": [], "headers": [], "table": None})

    return {"title": _txt(raw.get("title")) or L["rapport"], "sections": sections, "notes": notes}


# ======================================================================
# 4. TABLEAUX (valeurs issues de la BD)
# ======================================================================
_RISK_LABEL = {
    "rank": "rang", "ref": "ref", "domain": "domaine", "finding": "ecart", "severity": "sev",
    "freq": "freq", "class": "classe", "status": "conformite", "why": "pourquoi", "impact": "impact",
}
_DEFAULT_RISK_COLS = ["rank", "ref", "domain", "finding", "severity", "class", "why", "impact"]


def _risk_cell(key: str, i: int, it: dict, note: dict) -> str:
    if key == "rank":
        return str(i)
    if key == "ref":
        return _txt(it.get("item_id"))
    if key == "domain":
        sec, eq = _txt(it.get("section")), _txt(it.get("equipement"))
        return f"{sec} — {eq}" if sec and eq else (sec or eq)
    if key == "finding":
        return _txt(it.get("symptome"))
    if key == "severity":
        return _txt(it.get("severity"))
    if key == "freq":
        try:
            return f"{float(it.get('freq_moyenne') or 0):.1e}"
        except (TypeError, ValueError):
            return ""
    if key == "class":
        return _txt(it.get("risk_class"))
    if key == "status":
        return _txt(it.get("conformite"))
    if key == "why":
        return note.get("why_critical", "")
    if key == "impact":
        return note.get("potential_impact", "")
    return ""


def resolve_table(sec: dict, ctx: dict, L: dict):
    """-> (headers, rows) ou None. ctx = {items, stats, sections, notes}."""
    kind = sec.get("data_table")
    if kind == "priority_risks":
        cols = [c for c in sec.get("columns", []) if c in _RISK_LABEL] or _DEFAULT_RISK_COLS
        custom = sec.get("headers") or []
        headers = custom if len(custom) == len(cols) else [L[_RISK_LABEL[c]] for c in cols]
        rows = [[_risk_cell(c, i, it, ctx["notes"].get(str(it.get("item_id")), {})) for c in cols]
                for i, it in enumerate(ctx["items"], 1)]
        return headers, rows
    if kind == "compliance_summary":
        st = ctx["stats"]
        total = st.get("total_items", 0)
        n_c = st.get("n_c", 0)
        rate = (100.0 * n_c / total) if total else 0.0
        return ([L["repartition"], "#"],
                [[L["total"], str(total)], [L["c"], str(n_c)], [L["pc"], str(st.get("n_pc", 0))],
                 [L["nc"], str(st.get("n_nc", 0))], [L["conf_rate"], f"{rate:.0f}%"]])
    if kind == "domain_breakdown":
        return ([L["domaine"].split("/")[0].strip(), L["domaine_nc"], L["domaine_pc"], "C"],
                [[_txt(s["section"]), str(s["nc"]), str(s["pc"]), str(s["c"])] for s in ctx["sections"]])
    tbl = sec.get("table")
    if tbl:
        headers = [_txt(h) for h in _as_list(tbl.get("headers"))]
        rows = []
        for r in _as_list(tbl.get("rows")):
            cells = list(r.values()) if isinstance(r, dict) else _as_list(r)
            rows.append([_txt(c) for c in cells])
        if headers or rows:
            return headers, rows
    return None


_RISK_WEIGHT = {"rank": 5, "ref": 8, "domain": 16, "finding": 28, "severity": 10, "freq": 9,
                "class": 10, "status": 9, "why": 19, "impact": 19}


def column_weights(sec: dict, headers: list, rows: list) -> list[float]:
    """Poids relatifs des colonnes (largeur Word) : fixes pour le tableau des écarts,
    sinon proportionnels à la longueur du contenu."""
    ncols = max([len(headers)] + [len(r) for r in rows] + [1])
    if sec.get("data_table") == "priority_risks":
        cols = [c for c in sec.get("columns", []) if c in _RISK_LABEL] or _DEFAULT_RISK_COLS
        if len(cols) == ncols:
            return [float(_RISK_WEIGHT[c]) for c in cols]
    out = []
    for j in range(ncols):
        lens = [len(headers[j])] if j < len(headers) else []
        lens += [len(r[j]) for r in rows if j < len(r)]
        avg = sum(lens) / len(lens) if lens else 8
        out.append(min(max(avg, 6.0), 40.0))
    return out


def _ctx(report: dict, items: list, stats: dict, sections: list) -> dict:
    return {"items": items, "stats": stats, "sections": sections, "notes": report.get("notes", {})}


# ======================================================================
# 5. APERÇU HTML
# ======================================================================
_CUSTOM_CSS = """
h1.rt{margin:0 0 6px;font-size:26px;color:var(--navy)}
.rmeta{margin:10px 0 22px}
.rmeta td:first-child{width:30%;font-weight:700;background:#EEF3F8;color:var(--navy)}
.csec p{margin:6px 0}
.csec ul{margin:6px 0 6px 20px;padding:0}
.csec table{margin:10px 0 14px}
"""


def render_custom_report_html(report: dict, items: list, stats: dict, sections: list,
                              meta: dict, lang: str = "fr") -> str:
    L = _labels(lang)
    ctx = _ctx(report, items, stats, sections)
    site = meta.get("site", "—")
    tbc = L["a_completer"]
    logo = _logo_b64()

    meta_rows = [
        (L["site"], site), (L["inspection_id"], meta.get("inspection_id", "—")),
        (L["date"], meta.get("date") or datetime.now().strftime("%d/%m/%Y")),
        (L["inspecteur"], meta.get("inspector") or tbc), (L["verificateur"], meta.get("reviewer") or tbc),
        (L["version"], meta.get("version") or "1.0"),
    ]
    meta_html = "".join(f"<tr><td>{_x(k)}</td><td>{_x(v)}</td></tr>" for k, v in meta_rows)
    parts = [f"""<div class="page">{f'<img src="{logo}" alt="logo" style="height:54px;margin-bottom:12px">' if logo else ''}
      <h1 class="rt">{_x(report['title'])}</h1>
      <table class="ctl rmeta">{meta_html}</table>"""]

    for sec in report["sections"]:
        h = ""
        if sec["heading"]:
            if sec["level"] == 1:
                h = f'<h2 class="sec">{_x(sec["heading"])}</h2>'
            else:
                h = f'<h3 style="font-size:{15 if sec["level"] == 2 else 13.5}px">{_x(sec["heading"])}</h3>'
        paras = "".join(f"<p>{_x(p)}</p>" for p in sec["paragraphs"])
        bullets = f"<ul>{''.join(f'<li>{_x(b)}</li>' for b in sec['bullets'])}</ul>" if sec["bullets"] else ""
        tb = resolve_table(sec, ctx, L)
        tbl = ""
        if tb:
            headers, rows = tb
            ths = "".join(f"<th>{_x(c)}</th>" for c in headers)
            trs = "".join("<tr class=\"nobreak\">" + "".join(f"<td>{_x(c)}</td>" for c in r) + "</tr>" for r in rows)
            tbl = f"<table>{f'<tr>{ths}</tr>' if ths else ''}{trs}</table>"
        parts.append(f'<div class="csec">{h}{paras}{bullets}{tbl}</div>')

    parts.append(f'<p class="muted" style="font-size:11.5px;margin-top:22px">{_x(L["gen_note"])}</p></div>')
    footer = f"{L['confidentiel']} · {site}"
    return _page("".join(parts), report["title"], lang, footer, _REPORT_CSS + _CUSTOM_CSS)


# ======================================================================
# 6. EXPORT WORD (.docx) — construit sur le modèle de l'entreprise si .docx
# ======================================================================
_PH_RE = re.compile(r"\{\{\s*(\w+)\s*\}\}")


def _iter_paragraphs(container):
    """Tous les paragraphes d'un conteneur (corps / en-tête / cellule), tableaux imbriqués compris."""
    for p in container.paragraphs:
        yield p
    for t in getattr(container, "tables", []):
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _fill_placeholders(doc, values: dict) -> None:
    def fix(par):
        txt = par.text
        if "{{" not in txt:
            return
        new = _PH_RE.sub(lambda m: values.get(m.group(1).lower(), m.group(0)), txt)
        if new != txt and par.runs:
            par.runs[0].text = new
            for r in par.runs[1:]:
                r.text = ""

    for par in _iter_paragraphs(doc):
        fix(par)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer,
                     sec.even_page_header, sec.even_page_footer):
            try:
                for par in _iter_paragraphs(part):
                    fix(par)
            except Exception:
                continue


def build_docx(report: dict, items: list, stats: dict, sections: list, meta: dict,
               lang: str = "fr", template_bytes: bytes | None = None) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    L = _labels(lang)
    rtl = L["dir"] == "rtl"
    ctx = _ctx(report, items, stats, sections)
    tbc = L["a_completer"]

    # --- Document de base : le modèle .docx de l'entreprise, ou un document standard ---
    table_style = None
    if template_bytes:
        doc = Document(io.BytesIO(template_bytes))
        for t in doc.tables:
            if t.style is not None and t.style.name:
                table_style = t.style.name
                break
    else:
        doc = Document()
        doc.styles["Normal"].font.size = Pt(10.5)
    body = doc.element.body

    marker = None
    if template_bytes:
        for p in doc.paragraphs:
            if BODY_MARKER in p.text:
                marker = p
                break
        if marker is None:  # pas de repère : le corps du modèle est remplacé
            for el in list(body):
                if el.tag != qn("w:sectPr"):
                    body.remove(el)
        else:  # repère : on supprime l'exemple de structure qui le suit (jusqu'à {{END_REPORT_BODY}} ou la fin)
            from docx.text.paragraph import Paragraph
            removing = False
            for el in list(body):
                if el is marker._p:
                    removing = True
                    continue
                if not removing:
                    continue
                if el.tag == qn("w:sectPr"):
                    break
                is_end = el.tag == qn("w:p") and END_MARKER in Paragraph(el, doc).text
                body.remove(el)
                if is_end:
                    break
    before = list(body)

    # --- helpers RTL ---
    def rtl_par(p):
        if not rtl:
            return
        pPr = p._p.get_or_add_pPr()
        pPr.insert_element_before(
            OxmlElement("w:bidi"), "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
            "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
            "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
            "w:rPr", "w:sectPr", "w:pPrChange")
        for r in p.runs:
            r.font.rtl = True

    def add_par(text, style=None, bold=False, size=None):
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold or None
        if size:
            r.font.size = Pt(size)
        rtl_par(p)
        return p

    def add_heading(text, level):
        for name in (f"Heading {level}", f"Titre {level}"):
            try:
                doc.styles[name]
                p = doc.add_paragraph(text, style=name)
                rtl_par(p)
                return
            except KeyError:
                continue
        add_par(text, bold=True, size={1: 16, 2: 13, 3: 11.5}[level])

    def add_bullet(text):
        for name in ("List Bullet", "Liste à puces"):
            try:
                doc.styles[name]
                return add_par(text, style=name)
            except KeyError:
                continue
        return add_par("• " + text)

    def add_table(headers, rows, weights=None):
        ncols = max([len(headers)] + [len(r) for r in rows] + [1])
        tbl = doc.add_table(rows=0, cols=ncols)
        styled = False
        for name in filter(None, (table_style, "Table Grid", "Grille du tableau")):
            try:
                tbl.style = name
                styled = True
                break
            except Exception:
                continue
        if not styled:  # bordures manuelles
            tblPr = tbl._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                b = OxmlElement(f"w:{edge}")
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "4")
                b.set(qn("w:color"), "999999")
                borders.append(b)
            tblPr.insert_element_before(
                borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                "w:tblCaption", "w:tblDescription", "w:tblPrChange")
        if rtl:
            tblPr = tbl._tbl.tblPr
            tblPr.insert_element_before(
                OxmlElement("w:bidiVisual"), "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
                "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                "w:tblCaption", "w:tblDescription", "w:tblPrChange")

        def fill_row(values, header=False):
            cells = tbl.add_row().cells
            for j in range(ncols):
                cell = cells[j]
                par = cell.paragraphs[0]
                run = par.add_run(values[j] if j < len(values) else "")
                run.font.size = Pt(9)
                if header:
                    run.bold = True
                    if not table_style:  # fond d'en-tête seulement si pas de style d'entreprise
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:color"), "auto")
                        shd.set(qn("w:fill"), "E9EEF4")
                        tcPr.append(shd)
                rtl_par(par)

        if headers:
            fill_row(headers, header=True)
        for r in rows:
            fill_row(r)
        if weights and len(weights) == ncols:  # largeurs de colonnes proportionnelles au contenu
            try:
                ps = doc.sections[-1]
                avail = ps.page_width - ps.left_margin - ps.right_margin
                tot = float(sum(weights))
                tbl.autofit = False
                for j, col in enumerate(tbl.columns):  # grille du tableau (lue par Word/LibreOffice)
                    col.width = int(avail * weights[j] / tot)
                for row in tbl.rows:
                    for j, cell in enumerate(row.cells):
                        cell.width = int(avail * weights[j] / tot)
            except Exception:
                pass  # largeur automatique si le modèle n'expose pas ses marges
        doc.add_paragraph()  # espace après le tableau

    # --- Contenu ---
    try:  # titre : style « Title » de l'entreprise s'il existe, sinon gras 20 pt
        doc.styles["Title"]
        rtl_par(doc.add_paragraph(report["title"], style="Title"))
    except KeyError:
        add_par(report["title"], bold=True, size=20)
    meta_rows = [
        (L["site"], meta.get("site", "—")), (L["inspection_id"], meta.get("inspection_id", "—")),
        (L["date"], meta.get("date") or datetime.now().strftime("%d/%m/%Y")),
        (L["inspecteur"], meta.get("inspector") or tbc), (L["verificateur"], meta.get("reviewer") or tbc),
        (L["version"], meta.get("version") or "1.0"),
    ]
    if not template_bytes:  # tableau d'identification : seulement avec la mise en page standard
        add_table([], [[k, str(v)] for k, v in meta_rows])

    for sec in report["sections"]:
        if sec["heading"]:
            add_heading(sec["heading"], sec["level"])
        for ptxt in sec["paragraphs"]:
            add_par(ptxt)
        for b in sec["bullets"]:
            add_bullet(b)
        tb = resolve_table(sec, ctx, L)
        if tb:
            add_table(tb[0], tb[1], column_weights(sec, tb[0], tb[1]))
    add_par(L["gen_note"], size=8.5)

    # --- Insertion à l'emplacement du repère {{REPORT_BODY}} ---
    if marker is not None:
        new_els = [e for e in body if e not in before]
        for e in new_els:
            marker._p.addprevious(e)
        body.remove(marker._p)

    # --- Variables {{site}}, {{date}}... (corps, en-têtes, pieds de page) ---
    ref_code = "HSE-RAP-" + re.sub(r"[^A-Za-z0-9]+", "-", str(meta.get("inspection_id", ""))).strip("-")[:24]
    _fill_placeholders(doc, {
        "site": str(meta.get("site", "")), "inspection_id": str(meta.get("inspection_id", "")),
        "date": str(meta.get("date") or datetime.now().strftime("%d/%m/%Y")),
        "inspector": str(meta.get("inspector") or ""), "reviewer": str(meta.get("reviewer") or ""),
        "version": str(meta.get("version") or "1.0"), "ref": ref_code, "title": report["title"],
        # Anciens noms (modèles d'entreprise créés avant le renommage audit -> inspection)
        "audit_id": str(meta.get("inspection_id", "")),
        "auditor": str(meta.get("inspector") or ""),
    })
    try:
        doc.core_properties.title = report["title"]
    except Exception:
        pass

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
