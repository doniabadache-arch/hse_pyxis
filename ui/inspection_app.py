"""
inspection_app.py (UNIFIED v13 — Ordre exact du cahier des charges)
=====================================================================
1. Profil Entreprise (Nom, Domaine, Éco, RH, Légal, Env)
2. Configuration Matrice
3. Programme d'inspection
4. Inspection
5. Analyse QRA
6. Résultats & Classification
7. Plan PDCA
8. Rapport Final
9. Tableau de bord + Actions Correctives
10. Comparaison
11. Administration
"""

import streamlit as st
import json
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

# ---- Modules internes ----
import ai_service
import email_service
from inspection_db import (
    InspectionDB,
    DOMAIN_ALL,
    item_is_active,
    item_section_id,
    split_domains,
)
from matrice_manager import MatriceManager
from risk_matrix import RiskMatrixConfig

from monte_carlo_engine import (
    simuler_frequence_triangulaire,
    resumer_distribution,
    N_ITER_APERCU,
    N_ITER_FINAL,
)
from human_risk_indicators import (
    evaluer_risque_humain,
    niveau_lira,
    SEUILS_LIRA_INDICATIFS,
)
from cost_benefit import (
    evaluer_cout_benefice,
    CURRENCIES,
    DEFAULT_CURRENCY,
    get_default_indisponibilite,
    convert_from_eur,
)
from dispersion_aloha import evaluer_dispersion
from company_profile import (
    build_company_profile_text,
    DEFAULT_PROFILE,
    BUDGET_LEVELS,
    normalize_budget_level,
)

from results_viewer import (
    load_matrix_dfs,
    classify_results_df,
    render_metric_cards,
    render_heatmap,
    render_plots,
    freq_to_prob_level,
    export_to_excel,       
    export_to_pdf,         
)

from ui.theme import apply_theme, current_theme
from ui.i18n import t, set_language, get_current_lang
from ui.display import (
    CONFORMITE_A_INSPECTER,
    CONFORMITE_EMOJI,
    conformite_display_options,
    conformite_from_display,
    tr_conformite,
    tr_lira,
    tr_matrix_label,
    tr_symptom,
    tr_section,
    tr_equipment,
    tr_method_of,
    symptom_search_text,
    tr_catalog_name,
)
from ui.layout import render_header, render_footer
from ui.pages.welcome import render_welcome
from ui.pages.cal_page import render_calendar_page

# ---- Authentification / rôles / journal d'audit ----
import auth
from ui.pages.login import render_account_box, render_read_only_notice
from ui.pages.admin import render_admin_page

# ---- Pages modulaires ----
from ui.pages.dashboard import render_dashboard_page
from ui.pages.comparison import render_comparison_page
from ui.pages.corrective_actions import render_corrective_actions_page

# ---- Exports & Services ----
from dispersion_wizard import render_dispersion_wizard
from ghs_service import render_ghs_block


# ============================================================
# CONFIGURATION STREAMLIT
# ============================================================
st.set_page_config(
    page_title=t("app.full_title"),
    page_icon="🛡️",
    layout="wide",
)

apply_theme()

if not st.session_state.get("welcome_done", False):
    render_welcome()
    st.stop()

render_header()

# 🔐 Vérification de la connexion
current_user = auth.require_login()


# ============================================================
# INITIALISATION
# ============================================================
db = auth.guard_db(InspectionDB())
matrice_manager = auth.guard_matrice(MatriceManager())


# ============================================================
# SESSION STATE
# ============================================================
if "current_step" not in st.session_state:
    st.session_state.current_step = 8  # 👈 Démarre sur Profil Entreprise

if "inspection_buffer" not in st.session_state:
    st.session_state.inspection_buffer = {}

if "batch_done" not in st.session_state:
    st.session_state.batch_done = False

if "active_site" not in st.session_state:
    sites = matrice_manager.list_sites()
    st.session_state.active_site = sites[0] if sites else None

if "company_profile_data" not in st.session_state:
    st.session_state.company_profile_data = dict(DEFAULT_PROFILE)

# ═══ الإيميل ═══
if "email_recipient" not in st.session_state:
    st.session_state.email_recipient = ""
if "email_last_result" not in st.session_state:
    st.session_state.email_last_result = None


# ============================================================
# SIDEBAR — NAVIGATION
# ============================================================
render_account_box()

st.sidebar.header(f"🧭 {t('nav.navigation')}")

STEP_COMPANY = 8  # 👈 Profil Entreprise = 8
STEP_DASHBOARD = auth.STEP_DASHBOARD
STEP_COMPARISON = auth.STEP_COMPARISON
STEP_ACTIONS_CORR = auth.STEP_ACTIONS_CORR

# ═══ ORDRE LOGIQUE (Profil en 1er, Admin en dernier) ═══
_ALL_STEPS = {
    STEP_COMPANY: t("nav.company_profile_step"),  # 1. Profil Entreprise
    0: t("nav.step_0"),                           # 2. Configuration Matrice
    6: t("nav.step_1"),                           # 3. Programme d'inspection
    1: t("nav.step_2"),                           # 4. Inspection (Saisie)
    2: t("nav.step_3"),                           # 5. Analyse QRA
    3: t("nav.step_4"),                           # 6. Résultats & Classification
    4: t("nav.step_5"),                           # 7. Plan PDCA
    5: t("nav.step_6"),                           # 8. Rapport Final
    STEP_DASHBOARD: t("nav.dashboard"),           # 9. Dashboard & Actions
    STEP_COMPARISON: t("nav.comparison"),         # 10. Comparaison
    auth.STEP_ADMIN: t("nav.step_7"),             # 11. Administration (7)
}

STEPS = {k: v for k, v in _ALL_STEPS.items() if auth.can_access_step(k)}

if st.session_state.current_step not in STEPS:
    st.session_state.current_step = next(iter(STEPS))
if st.session_state.get("navigation_radio") not in STEPS:
    st.session_state.pop("navigation_radio", None)


def _goto_step(n):
    st.session_state["current_step"] = n
    st.session_state["navigation_radio"] = n


def _cb_plan_followup(site, sections, notes):
    ss = st.session_state
    ss["cal_new_site"] = site
    ss["cal_new_sections"] = list(sections)
    ss["cal_new_notes"] = notes
    ss["cal_new_date"] = date.today() + timedelta(days=30)
    ss["calendar_open_form"] = True
    _goto_step(6)

step = st.sidebar.radio(
    t("nav.active_step"),
    options=list(STEPS.keys()),
    format_func=lambda x: STEPS[x],
    index=list(STEPS).index(st.session_state.current_step),
    key="navigation_radio",
)
st.session_state.current_step = step


# ============================================================
# SIDEBAR — INFOS D'INSPECTION
# ============================================================
st.sidebar.markdown("---")
st.sidebar.header(f"📋 {t('sidebar.inspection_info')}")

inspection_id = st.sidebar.text_input(
    t("sidebar.inspection_id"),
    value="INSPECTION-2025-001",
    key="inspection_id_input",
)

_saved_insp = db.get_inspection_date(inspection_id)
try:
    _insp_default = date.fromisoformat(_saved_insp) if _saved_insp else date.today()
except ValueError:
    _insp_default = date.today()
inspection_date = st.sidebar.date_input(
    t("sidebar.inspection_date"),
    value=_insp_default,
    format="DD/MM/YYYY",
    key=f"inspection_date_input_{inspection_id}",
    disabled=not auth.has_permission("edit_inspection"),
)
if (inspection_date and inspection_date.isoformat() != _saved_insp
        and auth.has_permission("edit_inspection")):
    db.save_inspection_date(inspection_id, inspection_date.isoformat())
st.session_state["inspection_date_str"] = (
    inspection_date.strftime("%d/%m/%Y") if inspection_date else ""
)


# ============================================================
# SIDEBAR — STATUT DU BUFFER
# ============================================================
st.sidebar.markdown("---")
st.sidebar.header(f"📊 {t('sidebar.inspection_status')}")

buffer = st.session_state.inspection_buffer
n_total = len(buffer)
n_nc = sum(1 for v in buffer.values() if v["conformite"] == "NC")
n_pc = sum(1 for v in buffer.values() if v["conformite"] == "PC")
n_c = sum(1 for v in buffer.values() if v["conformite"] == "C")

st.sidebar.metric(t("sidebar.total_entered"), n_total)

col1, col2 = st.sidebar.columns(2)
col1.metric(tr_conformite("NC"), n_nc)
col2.metric(tr_conformite("PC"), n_pc)
st.sidebar.metric(tr_conformite("C"), n_c)

if st.sidebar.button(t("sidebar.reset_buffer")):
    st.session_state.inspection_buffer = {}
    st.session_state.batch_done = False
    st.sidebar.success(t("sidebar.buffer_reset"))
    st.rerun()


# ============================================================
# SIDEBAR — PROCHAINES INSPECTIONS / ACTIONS EN RETARD
# ============================================================
st.sidebar.markdown("---")
st.sidebar.header(f"📅 {t('sidebar.programme')}")

_today_iso = date.today().isoformat()
_upcoming = [s_ for s_ in db.get_scheduled_inspections(date_from=_today_iso)
             if s_["status"] != "done"][:3]
if _upcoming:
    for s_ in _upcoming:
        st.sidebar.caption(
            f"• {s_['scheduled_date']} {s_.get('scheduled_time') or ''} — "
            f"{s_.get('site_name') or '—'}"
        )
else:
    st.sidebar.caption(t("sidebar.no_upcoming"))

_n_late = sum(1 for a_ in db.get_corrective_actions(status="open")
              if a_.get("deadline") and a_["deadline"] < _today_iso)
if _n_late:
    st.sidebar.warning(t("sidebar.actions", n=_n_late))
st.sidebar.button(t("sidebar.open_programme"), key="sidebar_open_programme",
                  on_click=_goto_step, args=(6,))


# ============================================================
# SIDEBAR — PARAMÈTRES EMAIL
# ============================================================
if auth.has_permission("send_email"):
    st.sidebar.markdown("---")
    st.sidebar.header("📧 Email")

    email_status = email_service.get_status()
    if not email_status["configured"]:
        st.sidebar.warning("⚠️ SMTP non configuré")
        st.sidebar.caption("Ajoutez SMTP_* dans `.env`")
    else:
        providers = email_status.get("smtp_providers", [])
        if email_status.get("api_configured"):
            st.sidebar.success("✅ Brevo API configurée")
        for i, p in enumerate(providers, 1):
            role = "Fallback" if p["is_fallback"] else "Principal"
            st.sidebar.caption(f"{i}. {role} : {p['label']} ({p['mode']})")

        st.session_state.email_recipient = st.sidebar.text_input(
            "📬 Destinataire",
            value=st.session_state.email_recipient,
            placeholder="manager@company.com",
            key="email_recipient_input",
        )

        if st.sidebar.button("📧 Tester l'envoi", key="test_email_btn"):
            if not st.session_state.email_recipient:
                st.sidebar.error("Destinataire vide")
            else:
                with st.sidebar.spinner("Test..."):
                    r = email_service.send_email(
                        to=st.session_state.email_recipient,
                        subject="🎉 HSE PYXIS — Test Email",
                        html_body="""
                        <html><body style="font-family:Arial;">
                        <h1 style="color:#00C853;">🎉 Test réussi !</h1>
                        <p>Votre configuration SMTP fonctionne.</p>
                        </body></html>
                        """,
                    )
                if r["ok"]:
                    st.sidebar.success(f"✅ Envoyé via {r['provider']}")
                else:
                    st.sidebar.error(f"❌ {r['error']}")


# ============================================================
# SIDEBAR — PROFIL ENTREPRISE (résumé ; le formulaire complet est l'étape « Profil entreprise »)
# ============================================================
st.sidebar.markdown("---")
st.sidebar.header(f"🏢 {t('sidebar.company_profile')}")

_sb_profile = st.session_state.company_profile_data
_sb_dom = _sb_profile.get("domaine_activite", "generique")
st.sidebar.markdown(
    f"**{_sb_profile.get('nom_entreprise') or '—'}**  \n"
    f"{t('company.domain_label')} : {t(f'company.domain_{_sb_dom}')}"
)
if st.session_state.get("company_profile_saved"):
    st.sidebar.success(t(
        "sidebar.company_profile_summary",
        size=t(f"company.size_{_sb_profile.get('taille', 'PME')}"),
        budget=t(f"company.budget_hse_{normalize_budget_level(_sb_profile.get('budget_hse'), 'moyen')}").split(" — ")[0],
    ))
if STEP_COMPANY in STEPS:
    st.sidebar.button(
        t("company.edit"), key="cp_edit",
        on_click=_goto_step, args=(STEP_COMPANY,),
    )


# ============================================================
# SIDEBAR — PARAMÈTRES (Langue + Thème)
# ============================================================
st.sidebar.markdown("---")
st.sidebar.header(f"⚙️ {t('nav.settings')}")

from ui.i18n import LANGUAGES
lang_options = list(LANGUAGES.keys())
current_lang = get_current_lang()
lang_idx = lang_options.index(current_lang) if current_lang in lang_options else 0

new_lang = st.sidebar.selectbox(
    f"🌍 {t('nav.language')}",
    lang_options,
    index=lang_idx,
    format_func=lambda x: f"{LANGUAGES[x]['flag']} {LANGUAGES[x]['label']}",
    key="sidebar_lang_select",
)
if new_lang != current_lang:
    set_language(new_lang)
    st.rerun()

theme_actuel = current_theme()
new_theme = st.sidebar.radio(
    f"🎨 {t('nav.theme')}",
    ["dark", "light"],
    index=0 if theme_actuel == "dark" else 1,
    format_func=lambda x: t("welcome.theme_dark") if x == "dark" else t("welcome.theme_light"),
    key="sidebar_theme_radio",
)
if new_theme != theme_actuel:
    st.session_state.theme = new_theme
    st.rerun()


# ============================================================
# LOAD SUBSTANCES
# ============================================================
@st.cache_data
def load_substances():
    f = Path("data/substances_clean.json")
    if not f.exists():
        return []
    with open(f, "r", encoding="utf-8") as fp:
        return json.load(fp)


substances_db = load_substances()
substance_options = ["-- Sélectionner --"] + [
    s.get("name") or f"Unknown ({s.get('source_filename')})"
    for s in substances_db
]


# ============================================================
# LOAD CHECKLIST DEPUIS EXCEL
# ============================================================
def _xl_int(value, default=None):
    """int(value) tolérant (1, 1.0, '1') ; `default` si vide / invalide."""
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _xl_bool(value, default=True):
    """TRUE/FALSE Excel (bool, 1/0 ou texte) ; `default` si la cellule est vide."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().upper() in ("TRUE", "1", "YES", "OUI")


def load_checklist_from_excel():
    import openpyxl

    xlsx_path = Path("anomaly checklist/HSE_Checklist_Final_118.xlsx")

    if not xlsx_path.exists():
        st.error(t("checklist.file_not_found", path=xlsx_path.resolve()))
        st.info(t("checklist.check_folder"))
        return []

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        # Feuille des symptômes : « Sheet1 » (le classeur contient aussi Domains / Sections)
        ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]

        headers = [cell.value for cell in ws[1]]
        col_index = {}
        for i, h in enumerate(headers):
            if h is not None:
                col_index[str(h).strip().lower()] = i

        with st.expander(t("checklist.debug_title")):
            st.write(t("checklist.debug_raw_headers"), headers)
            st.write(t("checklist.debug_indexed_columns"), list(col_index.keys()))

        def get_val(row_values, *names, default=None):
            for n in names:
                n_lower = n.lower().strip()
                if n_lower in col_index:
                    val = row_values[col_index[n_lower]]
                    if val is not None:
                        return val
            return default

        items = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            item_id = get_val(row, "id")
            if not item_id:
                continue

            try:
                freq_val = get_val(row, "frequency/year", "frequency", default=0.001)
                try:
                    freq_val = float(freq_val)
                except (ValueError, TypeError):
                    freq_val = 0.001

                sev_val = get_val(row, "severity", default=3)
                try:
                    sev_val = int(float(sev_val))
                except (ValueError, TypeError):
                    sev_val = 3

                def to_bool(val):
                    if isinstance(val, bool):
                        return val
                    return str(val).strip().upper() == "TRUE"

                items.append({
                    "id": str(item_id),
                    "section_name": str(get_val(row, "section_name", "section", default="")),
                    "equipement": str(get_val(row, "equipement_type", "equipement", default="")),
                    "section_name_fr": str(get_val(row, "section_name_fr", default="")).strip(),
                    "section_name_ar": str(get_val(row, "section_name_ar", default="")).strip(),
                    "equipement_fr": str(get_val(row, "equipement_type_fr", "equipement_fr", default="")).strip(),
                    "equipement_ar": str(get_val(row, "equipement_type_ar", "equipement_ar", default="")).strip(),
                    "method_fr": str(get_val(row, "frequency_method_fr", "method_fr", default="")).strip(),
                    "method_ar": str(get_val(row, "frequency_method_ar", "method_ar", default="")).strip(),
                    "hazard_class": str(get_val(row, "hazzard class", "hazzard_class", "hazard class", "hazard_class", default="")).strip(),
                    "hazard_class_fr": str(get_val(row, "hazzard_class_fr", "hazard_class_fr", default="")).strip(),
                    "hazard_class_ar": str(get_val(row, "hazzard_class_ar", "hazard_class_ar", default="")).strip(),
                    "symptome": str(get_val(row, "symptom", "symptome", default="")),
                    "symptome_fr": str(get_val(row, "symptome_fr", default="")).strip(),
                    "symptome_ar": str(get_val(row, "symptome_ar", default="")).strip(),
                    "severity": sev_val,
                    "frequency": freq_val,
                    "method": str(get_val(row, "frequency_method", "method", default="Estimated")),
                    "source": str(get_val(row, "frequency_source", "source", default="")),
                    "evidence": str(get_val(row, "evidence", default="")),
                    "target_human": to_bool(get_val(row, "target human", "target_human", default=False)),
                    "target_env": to_bool(get_val(row, "target environment", "target_env", default=False)),
                    "target_asset": to_bool(get_val(row, "target asset", "target_asset", default=False)),
                    # --- Domaine / section / actif (étapes 1-3 du classeur) ---
                    "sec": _xl_int(get_val(row, "sec")),
                    "domain_code": str(get_val(row, "domain_code", default="")).strip(),
                    "symptom_code": str(get_val(row, "symptom_code", default="")).strip(),
                    "is_active": _xl_bool(get_val(row, "is_active", default=True)),
                })
            except Exception as e:
                st.warning(t("checklist.row_skipped", row=row_idx, error=e))
                continue

        return items

    except Exception as e:
        st.error(t("checklist.excel_error", error=e))
        import traceback
        st.code(traceback.format_exc())
        return []


def load_catalog_from_excel():
    """Lit les feuilles « Domains » et « Sections » du classeur de la checklist.

    Retourne (domains, sections) : deux listes de dict ; ([], []) si les feuilles
    sont absentes (ancien classeur -> l'inspection affiche alors toute la checklist).
    """
    import openpyxl

    xlsx_path = Path("anomaly checklist/HSE_Checklist_Final_118.xlsx")
    if not xlsx_path.exists():
        return [], []
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        st.error(t("checklist.excel_error", error=e))
        return [], []

    def _records(sheet_name):
        if sheet_name not in wb.sheetnames:
            return []
        ws_ = wb[sheet_name]
        heads = [str(c.value).strip().lower() if c.value is not None else "" for c in ws_[1]]
        out = []
        for row in ws_.iter_rows(min_row=2, values_only=True):
            rec = {h: v for h, v in zip(heads, row) if h}
            if any(v not in (None, "") for v in rec.values()):
                out.append(rec)
        return out

    def _txt(value):
        return "" if value is None else str(value).strip()

    domains = []
    for r in _records("Domains"):
        code = _txt(r.get("domain_code"))
        if not code or code.upper() == DOMAIN_ALL:   # ALL = valeur technique, pas un domaine
            continue
        domains.append({
            "domain_code": code,
            "name_en": _txt(r.get("name_en")),
            "name_fr": _txt(r.get("name_fr")),
            "name_ar": _txt(r.get("name_ar")),
            "icon": _txt(r.get("icon")),
            "display_order": _xl_int(r.get("display_order"), default=len(domains) + 1),
        })

    sections = []
    for r in _records("Sections"):
        sid = _xl_int(r.get("section_id"))
        if sid is None:
            continue
        sections.append({
            "section_id": sid,
            "name_en": _txt(r.get("section_name_en")),
            "name_fr": _txt(r.get("section_name_fr")),
            "name_ar": _txt(r.get("section_name_ar")),
            "domain_code": _txt(r.get("domain_code")),
            "display_order": _xl_int(r.get("display_order"), default=0),
            "is_active": _xl_bool(r.get("is_active"), default=True),
        })
    return domains, sections


def _apply_section_names(excel_items, sections):
    """Complète section_name_fr / section_name_ar des symptômes avec les noms de la
    feuille Sections (la feuille Sheet1 ne contient que le nom anglais)."""
    names = {s["section_id"]: s for s in sections}
    for it in excel_items:
        sec_row = names.get(it.get("sec"))
        if not sec_row:
            continue
        if not it.get("section_name_fr"):
            it["section_name_fr"] = sec_row.get("name_fr", "")
        if not it.get("section_name_ar"):
            it["section_name_ar"] = sec_row.get("name_ar", "")


items = db.get_checklist_items()

_refresh_translations = (
    bool(items)
    and not st.session_state.get("_checklist_i18n_checked", False)
    and (db.checklist_needs_translation_refresh() or db.checklist_needs_catalog_refresh())
)
st.session_state["_checklist_i18n_checked"] = True

if not items or _refresh_translations:
    with st.spinner(t("checklist.loading")):
        excel_items = load_checklist_from_excel()
        if excel_items:
            _domains, _sections = load_catalog_from_excel()
            _apply_section_names(excel_items, _sections)
            db.save_checklist_items(excel_items)
            db.save_domains(_domains)
            db.save_sections(_sections)
            items = db.get_checklist_items()
            if _refresh_translations:
                st.success(t("checklist.i18n_refreshed", n=len(items)))
            else:
                st.success(t("checklist.items_loaded", n=len(items)))
        elif not items:
            st.error(t("checklist.load_failed"))
            st.stop()


# ============================================================
# ÉTAPE 8 — PROFIL ENTREPRISE (Company Profile)
# ============================================================
def render_step_company_profile():
    """Étape 8 : Profil de l'entreprise (avant la configuration matrice)."""
    st.header(t("sidebar.company_profile"))
    st.caption(t("sidebar.company_profile_caption"))

    profile = st.session_state.company_profile_data

    tab1, tab2, tab3, tab4 = st.tabs([
        t("company.tab_economic"),
        t("company.tab_hr"),
        t("company.tab_legal"),
        t("company.tab_env"),
    ])

    with tab1:
        profile["nom_entreprise"] = st.text_input(
            "Nom de l'entreprise",
            value=profile.get("nom_entreprise", ""),
            key="cp_nom_ent_main",
        )
        profile["domaine_activite"] = st.selectbox(
            "Domaine d'activité",
            ["petrole_gaz", "mines", "construction", "manufacturing", "pharmacetique", "stockage_transport", "generique"],
            index=["petrole_gaz", "mines", "construction", "manufacturing", "pharmacetique", "stockage_transport", "generique"].index(
                profile.get("domaine_activite", "generique")),
            format_func=lambda x: {
                "petrole_gaz": "Pétrole & Gaz",
                "mines": "Mines",
                "construction": "Construction",
                "manufacturing": "Manufacturing",
                "pharmacetique": "Pharmaceutique",
                "stockage_transport": "Stockage & Transport",
                "generique": "Générique",
            }.get(x, x),
            key="cp_dom_act_main",
        )
        st.caption(t("company.domain_sections_hint"))
        profile["taille"] = st.selectbox(
            t("company.size"),
            ["artisan", "PME", "ETI", "grande", "multinationale"],
            index=["artisan", "PME", "ETI", "grande", "multinationale"].index(
                profile.get("taille", "PME")),
            format_func=lambda x: t(f"company.size_{x}"),
            key="cp_taille_main",
        )
        profile["budget_hse"] = st.selectbox(
            t("company.budget_hse_level"),
            BUDGET_LEVELS,
            index=BUDGET_LEVELS.index(
                normalize_budget_level(profile.get("budget_hse"), "moyen")),
            format_func=lambda x: t(f"company.budget_hse_{x}"),
            key="cp_budget_hse_level_main",
        )
        profile["priorite"] = st.selectbox(
            t("company.priority"),
            ["securite", "conformite", "production", "cout"],
            index=["securite", "conformite", "production", "cout"].index(
                profile.get("priorite", "securite")),
            format_func=lambda x: t(f"company.priority_{x}"),
            key="cp_priorite_main",
        )
        profile["delai"] = st.selectbox(
            t("company.timeline"),
            ["immediat", "1_mois", "3_mois", "1_an", "long_terme"],
            index=["immediat", "1_mois", "3_mois", "1_an", "long_terme"].index(
                profile.get("delai", "3_mois")),
            format_func=lambda x: t(f"company.timeline_{x}"),
            key="cp_delai_main",
        )

    with tab2:
        profile["niveau_formation"] = st.selectbox(
            t("company.training_level"),
            ["faible", "moyen", "eleve", "expert"],
            index=["faible", "moyen", "eleve", "expert"].index(
                profile.get("niveau_formation", "moyen")),
            format_func=lambda x: t(f"company.training_level_{x}"),
            key="cp_niv_formation_main",
        )
        profile["turnover"] = st.selectbox(
            t("company.turnover"),
            ["faible", "moyen", "eleve"],
            index=["faible", "moyen", "eleve"].index(
                profile.get("turnover", "moyen")),
            format_func=lambda x: t(f"company.turnover_{x}"),
            key="cp_turnover_main",
        )
        profile["formation_disponible"] = st.selectbox(
            t("company.training_availability"),
            ["aucune", "occasionnelle", "reguliere"],
            index=["aucune", "occasionnelle", "reguliere"].index(
                profile.get("formation_disponible", "occasionnelle")),
            format_func=lambda x: t(f"company.training_availability_{x}"),
            key="cp_form_dispo_main",
        )
        profile["budget_formation"] = st.selectbox(
            t("company.training_budget_level"),
            BUDGET_LEVELS,
            index=BUDGET_LEVELS.index(
                normalize_budget_level(profile.get("budget_formation"), "faible")),
            format_func=lambda x: t(f"company.budget_formation_{x}"),
            key="cp_budget_formation_level_main",
        )
        profile["metiers_specifiques"] = st.text_area(
            t("company.specific_jobs"),
            value=profile.get("metiers_specifiques", ""),
            height=60,
            placeholder=t("company.specific_jobs_placeholder"),
            key="cp_metiers_main",
        )

    with tab3:
        profile["pays"] = st.text_input(
            t("company.country"),
            value=profile.get("pays", "Algerie"),
            key="cp_pays_main",
        )
        profile["reglementation_applicable"] = st.text_area(
            t("company.regulation"),
            value=profile.get("reglementation_applicable", ""),
            height=60,
            key="cp_reglementation_main",
        )
        profile["normes_certifiees"] = st.text_input(
            t("company.certified_standards"),
            value=profile.get("normes_certifiees", ""),
            placeholder=t("company.certified_standards_placeholder"),
            key="cp_normes_cert_main",
        )
        profile["certifications_visees"] = st.text_input(
            t("company.target_certifications"),
            value=profile.get("certifications_visees", ""),
            key="cp_normes_visees_main",
        )
        profile["autorisations_requises"] = st.text_area(
            t("company.required_authorizations"),
            value=profile.get("autorisations_requises", ""),
            height=60,
            placeholder=t("company.required_authorizations_placeholder"),
            key="cp_autorisations_main",
        )
        profile["conventions_collectives"] = st.text_area(
            t("company.collective_agreements"),
            value=profile.get("conventions_collectives", ""),
            height=60,
            key="cp_conventions_main",
        )
        profile["secteur_reglemente"] = st.checkbox(
            t("company.regulated_sector"),
            value=profile.get("secteur_reglemente", False),
            key="cp_secteur_reg_main",
        )

    with tab4:
        profile["zone_habitee_proche"] = st.checkbox(
            t("company.nearby_residential"),
            value=profile.get("zone_habitee_proche", False),
            key="cp_zone_hab_main",
        )
        profile["zone_protegee_proche"] = st.checkbox(
            t("company.nearby_protected"),
            value=profile.get("zone_protegee_proche", False),
            key="cp_zone_prot_main",
        )
        profile["niveau_sonore_actuel"] = st.selectbox(
            t("company.current_noise"),
            ["faible", "moyen", "eleve"],
            index=["faible", "moyen", "eleve"].index(
                profile.get("niveau_sonore_actuel", "moyen")),
            format_func=lambda x: t(f"company.current_noise_{x}"),
            key="cp_bruit_main",
        )
        profile["emissions_actuelles"] = st.text_area(
            t("company.current_emissions"),
            value=profile.get("emissions_actuelles", ""),
            height=60,
            placeholder=t("company.current_emissions_placeholder"),
            key="cp_emissions_main",
        )
        profile["gestion_dechets"] = st.selectbox(
            t("company.waste_management"),
            ["standard", "renforcee", "a_ameliorer"],
            index=["standard", "renforcee", "a_ameliorer"].index(
                profile.get("gestion_dechets", "standard")),
            format_func=lambda x: t(f"company.waste_management_{x}"),
            key="cp_dechets_main",
        )
        profile["contraintes_env_particulieres"] = st.text_area(
            t("company.env_constraints"),
            value=profile.get("contraintes_env_particulieres", ""),
            height=60,
            key="cp_env_contraintes_main",
        )

    st.markdown("---")
    profile["notes"] = st.text_area(
        t("company.notes"),
        value=profile.get("notes", ""),
        height=60,
        key="cp_notes_main",
    )

    if st.button(t("company.save"), key="cp_save_main", type="primary", use_container_width=True):
        st.session_state.company_profile_data = profile
        st.session_state.company_profile_saved = True
        st.success(t("company.saved"))
        # Navigation automatique vers la Matrice
        st.session_state.current_step = 0
        st.session_state["navigation_radio"] = 0
        st.rerun()


# ============================================================
# ÉTAPE 0 — CONFIGURATION MATRICE
# ============================================================
COULEURS = ["#2ECC71", "#F1C40F", "#E67E22", "#E74C3C", "#8B0000", "#5B2C6F"]
_ACTION_KEYS = ["none", "monitoring", "planned", "priority", "immediate", "urgent"]
_SEUIL_KEYS = ["acceptable", "tolerable", "high", "critical", "unacceptable"]


def actions_par_defaut():
    return [t(f"matrix_defaults.action_{k}") for k in _ACTION_KEYS]


def labels_seuils_par_defaut():
    return [t(f"matrix_defaults.threshold_{k}") for k in _SEUIL_KEYS]


_STEP0_LABEL_KEY_PREFIXES = ("sev_label_", "sev_action_", "prob_label_", "seuil_label_")


def _is_default_matrix_label(value) -> bool:
    raw = str(value or "").strip()
    if not raw:
        return False
    return any(tr_matrix_label(raw, lang=code) != raw for code in LANGUAGES)


def _retranslate_default_matrix_labels():
    lang = get_current_lang()
    previous = st.session_state.get("_step0_labels_lang")
    st.session_state["_step0_labels_lang"] = lang
    if previous is None or previous == lang:
        return

    for key in list(st.session_state.keys()):
        if (
            isinstance(key, str)
            and key.startswith(_STEP0_LABEL_KEY_PREFIXES)
            and _is_default_matrix_label(st.session_state[key])
        ):
            del st.session_state[key]


def render_step_0_matrix_config():
    """Étape 0 : Configuration de la matrice de risque."""

    _retranslate_default_matrix_labels()

    st.header(t("step0.title"))
    st.caption(t("step0.caption"))

    # ---- CHOIX DU SITE ----
    sites = matrice_manager.list_sites()
    col1, col2 = st.columns([2, 2])
    choix = col1.selectbox(
        t("step0.site_label"),
        ["-- Nouveau site --"] + sites,
        format_func=lambda x: t("step0.new_site_option") if x == "-- Nouveau site --" else x,
        key="matrix_site_choice",
    )

    site_name = (
        col2.text_input(t("step0.new_site_name_label"), key="matrix_site_input")
        if choix == "-- Nouveau site --"
        else choix
    )
    site_name = site_name.strip() if site_name else ""

    if not site_name:
        st.info(t("step0.choose_or_new_site"))
        return

    if choix == "-- Nouveau site --" and site_name in sites:
        st.warning(t("step0.site_exists_warning"))

    st.markdown("---")

    # ---- NOMBRE DE NIVEAUX ----
    col1, col2 = st.columns(2)
    nb_severite = col1.number_input(
        t("step0.nb_severite_label"),
        min_value=2, max_value=6, value=5,
        key="nb_severite",
    )
    nb_probabilite = col2.number_input(
        t("step0.nb_probabilite_label"),
        min_value=2, max_value=6, value=5,
        key="nb_probabilite",
    )
    SCORE_MAX = nb_severite * nb_probabilite

    # ---- SÉVÉRITÉ ----
    st.header(t("step0.severity_header"))
    severite_existante = matrice_manager.get_matrix(site_name)

    severite_data = []
    for i in range(1, nb_severite + 1):
        ancien = next((row for row in severite_existante if row[0] == i), None)
        label_defaut = tr_matrix_label(ancien[1]) if ancien else t("step0.level_default", n=i)
        couleur_defaut = ancien[2] if ancien else COULEURS[(i - 1) % len(COULEURS)]
        _actions = actions_par_defaut()
        action_defaut = tr_matrix_label(ancien[3]) if ancien else _actions[(i - 1) % len(_actions)]

        col1, col2, col3 = st.columns([3, 2, 1])
        label = col1.text_input(t("step0.level_label", n=i), label_defaut, key=f"sev_label_{i}")
        action = col2.text_input(t("step0.action_label", n=i), action_defaut, key=f"sev_action_{i}")
        couleur = col3.color_picker(
            t("step0.color_label", n=i), couleur_defaut, key=f"sev_c_{i}",
            label_visibility="collapsed",
        )
        severite_data.append({"Level": i, "Label": label, "Action": action, "Color": couleur})

    # ---- PROBABILITÉ ----
    st.header(t("step0.probability_header"))
    st.caption(t("step0.probability_caption"))
    probabilite_existante = matrice_manager.get_probability(site_name)

    probabilite_data = []
    for i in range(1, nb_probabilite + 1):
        ancien = next((row for row in probabilite_existante if row[0] == i), None)
        label_defaut = tr_matrix_label(ancien[1]) if ancien else t("step0.level_default", n=i)
        freq_defaut = ancien[2] if ancien else ""
        couleur_defaut = ancien[3] if ancien else COULEURS[(i - 1) % len(COULEURS)]

        col1, col2, col3 = st.columns([2, 2, 1])
        label = col1.text_input(t("step0.level_label", n=i), label_defaut, key=f"prob_label_{i}")
        freq = col2.text_input(
            t("step0.frequency_label", n=i),
            freq_defaut, key=f"prob_freq_{i}",
        )
        couleur = col3.color_picker(
            t("step0.color_label", n=i), couleur_defaut, key=f"prob_c_{i}",
            label_visibility="collapsed",
        )
        probabilite_data.append({
            "Level": i, "Label": label,
            "FrequencyRange": freq, "Color": couleur,
        })

    # ---- SEUILS DE RISQUE ----
    st.header(t("step0.thresholds_header"))
    st.caption(t("step0.thresholds_caption", max=SCORE_MAX))
    seuils_existants = matrice_manager.get_thresholds(site_name)

    max_seuils = min(5, SCORE_MAX)
    nb_seuils = st.number_input(
        t("step0.nb_bands_label"),
        min_value=2, max_value=max_seuils,
        value=len(seuils_existants) if seuils_existants else 4,
        key="nb_seuils",
    )

    seuils_data = []
    for i in range(1, nb_seuils + 1):
        ancien = seuils_existants[i - 1] if len(seuils_existants) >= i else None
        score_defaut = ancien[0] if ancien else (SCORE_MAX * i) // nb_seuils
        _seuils_lbl = labels_seuils_par_defaut()
        label_defaut = tr_matrix_label(ancien[1]) if ancien else _seuils_lbl[(i - 1) % len(_seuils_lbl)]
        couleur_defaut = ancien[2] if ancien else COULEURS[(i - 1) % len(COULEURS)]

        col1, col2, col3 = st.columns([2, 2, 1])
        label = col1.text_input(t("step0.threshold_name_label", n=i), label_defaut, key=f"seuil_label_{i}")
        score = col2.number_input(
            t("step0.threshold_score_label", n=i),
            min_value=1, max_value=SCORE_MAX,
            value=min(score_defaut, SCORE_MAX),
            key=f"seuil_max_{i}",
        )
        couleur = col3.color_picker(
            t("step0.color_label", n=i), couleur_defaut, key=f"seuil_c_{i}",
            label_visibility="collapsed",
        )
        seuils_data.append({"MaxScore": score, "Label": label, "Color": couleur})

    seuils_data[-1]["MaxScore"] = SCORE_MAX
    seuils_data.sort(key=lambda row: row["MaxScore"])

    # ---- SAUVEGARDE ----
    st.markdown("---")
    if st.button(t("step0.save_button"), type="primary", key="save_matrix_btn"):
        labels_sev = [row["Label"] for row in severite_data]
        labels_prob = [row["Label"] for row in probabilite_data]
        scores = [row["MaxScore"] for row in seuils_data]

        if len(labels_sev) != len(set(labels_sev)) or len(labels_prob) != len(set(labels_prob)):
            st.error(t("step0.labels_unique_error"))
        elif scores != sorted(set(scores)):
            st.error(t("step0.thresholds_increasing_error"))
        else:
            try:
                matrice_manager.save_full_matrix(
                    site_name, severite_data, probabilite_data, seuils_data,
                )
                st.session_state.active_site = site_name
                st.success(t("step0.save_success", site=site_name))
            except Exception as e:
                st.error(t("step0.save_error", error=e))

    # ---- APERÇU ----
    st.header(t("step0.preview_header"))

    def niveau_de_risque(score, seuils):
        for seuil in seuils:
            if score <= seuil["MaxScore"]:
                return seuil["Label"]
        return seuils[-1]["Label"]

    tableau = []
    for prob in sorted(probabilite_data, key=lambda r: r["Level"], reverse=True):
        ligne = {t("step0.probability_column"): tr_matrix_label(prob["Label"])}
        for sev in severite_data:
            score = prob["Level"] * sev["Level"]
            niveau = tr_matrix_label(niveau_de_risque(score, seuils_data))
            ligne[tr_matrix_label(sev["Label"])] = f"{score} ({niveau})"
        tableau.append(ligne)

    st.dataframe(pd.DataFrame(tableau), use_container_width=True, hide_index=True)

    # ---- SUPPRESSION DE MATRICES ----
    st.markdown("---")
    st.header(t("step0.delete_header"))

    sites_existants = matrice_manager.list_sites()
    if not sites_existants:
        st.info(t("step0.no_matrix_to_delete"))
        return

    col1, col2 = st.columns([3, 1])
    with col1:
        site_a_supprimer = st.selectbox(
            t("step0.delete_selector_label"),
            sites_existants,
            key="delete_matrix_selector",
        )
    with col2:
        st.write("")
        st.write("")
        if st.button(t("step0.delete_button"), type="secondary", key="delete_matrix_btn"):
            st.session_state["confirm_delete"] = site_a_supprimer

    if st.session_state.get("confirm_delete"):
        site_conf = st.session_state["confirm_delete"]
        st.warning(t("step0.delete_confirm_warning", site=site_conf))

        col1, col2 = st.columns(2)
        with col1:
            if st.button(t("step0.delete_confirm_yes"), type="primary", key="confirm_delete_yes"):
                try:
                    matrice_manager.delete_matrix(site_conf)
                    st.success(t("step0.delete_success", site=site_conf))
                    st.session_state["confirm_delete"] = None
                    st.session_state["active_site"] = None
                    st.rerun()
                except Exception as e:
                    st.error(t("step0.delete_error", error=e))
        with col2:
            if st.button(t("step0.delete_confirm_no"), key="confirm_delete_no"):
                st.session_state["confirm_delete"] = None
                st.rerun()


# ============================================================
# DOMAINE (Profil entreprise) -> SECTIONS (Programme d'inspection) -> Inspection
# ============================================================
# « Domaine d'activité » du Profil entreprise -> code de domaine de la feuille Domains.
_PROFILE_DOMAIN_TO_CODE = {
    "petrole_gaz": "OIL_GAS",
    "mines": "MINES",
    "construction": "CONSTRUCTION",
    "manufacturing": "MANUFACTURING",
    "pharmacetique": "PHARMA",
    "stockage_transport": "TRANSPORT_STORAGE",
    "generique": None,          # pas de filtre : toutes les sections sont proposées
}


def _inspection_domain():
    """Domaine d'inspection = « Domaine d'activité » choisi dans le Profil entreprise.

    Retourne (domain_code | None, ligne du domaine | None). Quand le domaine change,
    les sections déjà choisies sont remises à zéro (elles appartiennent à l'ancien).
    """
    ss = st.session_state
    key = (ss.get("company_profile_data") or {}).get("domaine_activite", "generique")
    wanted = _PROFILE_DOMAIN_TO_CODE.get(key)
    row = next((d for d in db.get_domains() if d["domain_code"] == wanted), None) if wanted else None
    code = row["domain_code"] if row else None
    if code != ss.get("_insp_prev_domain", code):
        ss["insp_sections"] = []
    ss["_insp_prev_domain"] = code
    ss["insp_domain"] = code
    return code, row


def _domain_sections(dom_code):
    """Sections proposées : celles du domaine (+ communes « ALL »), ou toutes si générique."""
    return db.get_sections_for_domain(dom_code) if dom_code else db.get_sections()


def _items_by_section(all_items):
    out = {}
    for it in all_items:
        sid = item_section_id(it)
        if sid is not None and item_is_active(it):
            out.setdefault(sid, []).append(it)
    return out


def _selection_key(section_ids):
    """Clé courte et stable pour une sélection de sections (état du data_editor)."""
    import hashlib
    ids = sorted(str(s) for s in (section_ids or []))
    if not ids:
        return "all"
    return hashlib.md5("|".join(ids).encode("utf-8")).hexdigest()[:10]


def _render_section_selector(sections, items_by_sec, conformites, scope):
    """Sélecteur multiple de sections (+ Tout / Aucune / Non terminées).
    Mémorise le choix dans st.session_state['insp_sections'] et le retourne."""
    ss = st.session_state
    sec_by_id = {s["section_id"]: s for s in sections}
    sec_ids = list(sec_by_id)

    progress = {}
    for sid in sec_ids:
        its = items_by_sec.get(sid, [])
        done = sum(
            1 for it in its
            if conformites.get(it[0], CONFORMITE_A_INSPECTER) != CONFORMITE_A_INSPECTER
        )
        progress[sid] = (done, len(its))

    def _section_label(sid):
        s_row = sec_by_id[sid]
        done, total = progress[sid]
        shared = "🌐 " if DOMAIN_ALL in split_domains(s_row["domain_code"]) else ""
        return f"{shared}{sid} — {tr_catalog_name(s_row)}   ({done}/{total})"

    # Clé du widget propre au domaine ; initialisée (avant sa création) depuis le choix
    # mémorisé, en ignorant les sections qui ne sont plus valides.
    ms_key = f"insp_sections_select_{scope}"
    if ms_key not in ss:
        ss[ms_key] = [s for s in (ss.get("insp_sections") or []) if s in sec_by_id]

    def _set_sections(ids):
        ss[ms_key] = list(ids)

    unfinished_ids = [sid for sid in sec_ids if 0 < progress[sid][1] and progress[sid][0] < progress[sid][1]]
    b1, b2, b3, _sp = st.columns([1, 1, 1.4, 3])
    b1.button(t("step1.select_all_sections"), key=f"insp_sec_all_{scope}",
              on_click=_set_sections, args=(sec_ids,), use_container_width=True)
    b2.button(t("step1.clear_sections"), key=f"insp_sec_none_{scope}",
              on_click=_set_sections, args=([],), use_container_width=True)
    b3.button(t("step1.select_unfinished_sections"), key=f"insp_sec_todo_{scope}",
              on_click=_set_sections, args=(unfinished_ids,), use_container_width=True,
              disabled=not unfinished_ids)

    picked = st.multiselect(
        t("step1.section_label"),
        sec_ids,
        format_func=_section_label,
        placeholder=t("step1.section_placeholder"),
        key=ms_key,
    )
    # Ordre stable : celui d'affichage du domaine (et non l'ordre de clic).
    sel_secs = [sid for sid in sec_ids if sid in picked]
    ss["insp_sections"] = sel_secs
    return sel_secs


def _render_programme_section_picker(conformites):
    """Page « Programme d'inspection » : choisir les sections (du domaine du profil) à inspecter."""
    ss = st.session_state
    if not db.get_domains():
        return
    dom_code, dom_row = _inspection_domain()
    sections = _domain_sections(dom_code)

    st.subheader(f"🧩 {t('step1.programme_sections_title')}")
    if dom_row:
        c1, c2 = st.columns([4, 1])
        c1.markdown(
            f"**{t('step1.domain_label')}** : {dom_row.get('icon') or ''} "
            f"{tr_catalog_name(dom_row)}  \n_{t('step1.domain_from_profile')}_"
        )
        if STEP_COMPANY in STEPS:
            c2.button(t("step1.change_domain"), key="prog_change_domain",
                      on_click=_goto_step, args=(STEP_COMPANY,), use_container_width=True)
    else:
        st.info(t("step1.domain_generic_note"))
    if not sections:
        st.warning(t("step1.no_sections_for_domain"))
        return

    all_items = db.get_checklist_items()
    items_by_sec = _items_by_section(all_items)
    sel_secs = _render_section_selector(sections, items_by_sec, conformites, dom_code or "ALL")
    if not sel_secs:
        st.info(t("step1.pick_section_hint"))
        return

    sec_by_id = {s["section_id"]: s for s in sections}
    n_items = sum(len(items_by_sec.get(sid, [])) for sid in sel_secs)
    st.success(t("step1.selected_summary", n_sections=len(sel_secs), n_items=n_items))
    st.caption(" • ".join(f"{sid} — {tr_catalog_name(sec_by_id[sid])}" for sid in sel_secs))

    names = []
    for sid in sel_secs:
        its = items_by_sec.get(sid)
        if its:
            names.append(tr_section(its[0]))
    b1, b2, _sp = st.columns([1.4, 1.4, 3])
    b1.button(t("step1.go_inspection"), key="prog_go_inspection", type="primary",
              on_click=_goto_step, args=(1,), use_container_width=True)
    if auth.has_permission("edit_schedule"):
        b2.button(t("step1.plan_with_sections"), key="prog_plan_with_sections",
                  on_click=_cb_plan_followup,
                  args=(ss.get("active_site"), names, ""), use_container_width=True)


def _pick_domain_and_sections(all_items, conformites):
    """Étape Inspection : le domaine vient du Profil entreprise et les sections du
    Programme d'inspection. Retourne les symptômes actifs des sections choisies
    (groupés par section), ou None tant qu'aucune section n'est choisie.
    Ancien classeur sans feuilles Domains / Sections : tous les symptômes actifs.
    """
    ss = st.session_state
    if not db.get_domains():
        ss["insp_sections"] = []
        return [it for it in all_items if item_is_active(it)]

    dom_code, dom_row = _inspection_domain()
    sections = _domain_sections(dom_code)
    sec_by_id = {s["section_id"]: s for s in sections}
    sel_secs = [sid for sid in sec_by_id if sid in (ss.get("insp_sections") or [])]
    ss["insp_sections"] = sel_secs

    if not sel_secs:
        st.info(t("step1.programme_pick_hint"))
        st.button(t("step1.go_programme"), key="insp_go_programme", type="primary",
                  on_click=_goto_step, args=(6,))
        return None

    items_by_sec = _items_by_section(all_items)
    section_items = [it for sid in sel_secs for it in items_by_sec.get(sid, [])]
    head = f"{dom_row.get('icon') or ''} {tr_catalog_name(dom_row)} › " if dom_row else ""
    c1, c2 = st.columns([4, 1])
    c1.subheader(
        f"{head}{t('step1.selected_summary', n_sections=len(sel_secs), n_items=len(section_items))}".strip()
    )
    c2.button(t("step1.edit_sections"), key="insp_edit_sections",
              on_click=_goto_step, args=(6,), use_container_width=True)
    st.caption(" • ".join(f"{sid} — {tr_catalog_name(sec_by_id[sid])}" for sid in sel_secs))
    if not section_items:
        st.warning(t("step1.no_items_in_section"))
        return None
    return section_items


# ============================================================
# ÉTAPE 1 — INSPECTION (Saisie)
# ============================================================
def render_step_1_inspection():
    """Étape 1 : Tableau d'inspection — tous les symptômes + conformité + evidence."""

    st.header(t("step1.title"))
    st.caption(t("step1.caption"))

    available_sites = matrice_manager.list_sites()
    if not available_sites:
        st.error(t("step1.no_matrix"))
        st.info(t("step1.goto_step0"))
        return

    st.subheader(t("step1.active_matrix"))

    default_site = st.session_state.get("active_site", available_sites[0])
    if default_site not in available_sites:
        default_site = available_sites[0]

    site_name = st.selectbox(
        t("step1.choose_matrix"),
        available_sites,
        index=available_sites.index(default_site),
        key="inspection_site_selector",
    )
    st.session_state.active_site = site_name

    try:
        matrix_config = RiskMatrixConfig(site_name=site_name)
        st.success(t(
            "step1.matrix_loaded",
            site=site_name,
            n_sev=len(matrix_config.severity_levels),
            n_prob=len(matrix_config.probability_levels),
        ))
    except Exception as e:
        st.error(t("common.matrix_error", error=e))
        return

    st.markdown("---")

    all_items = db.get_checklist_items()
    conformites_existantes = db.get_conformity_by_inspection(inspection_id)

    # Domaine -> Section(s) : `all_items` devient la liste des symptômes des sections choisies
    all_items = _pick_domain_and_sections(all_items, conformites_existantes)
    if all_items is None:
        return

    st.subheader(t("step1.inspection_table"))

    ai_verifications = db.get_all_ai_verifications(inspection_id)

    _ev_map = {}
    for _e in db.get_inspection_data(inspection_id):
        _ev_map[_e[3]] = {
            "text": (_e[25] if len(_e) > 25 else "") or "",
            "file": (_e[26] if len(_e) > 26 else None),
        }
    AI_EMOJI = {
        "conforme": "✅",
        "partiel": "⚠️",
        "non_conforme": "❌",
        "absent": "❓",
    }

    rows = []
    for item in all_items:
        item_id = item[0]
        section_name = tr_section(item)
        equipement = tr_equipment(item)
        severity = item[4]
        evidence = item[8] if len(item) > 8 else ""

        conformite_actuelle = conformites_existantes.get(item_id, CONFORMITE_A_INSPECTER)

        ai_data = ai_verifications.get(item_id)
        if ai_data:
            ai_emoji = AI_EMOJI.get(ai_data.get("verdict", ""), "❓")
            ai_status = f"{ai_emoji} {ai_data.get('score', 0)}"
        else:
            ai_status = "—"

        _ev = _ev_map.get(item_id, {})
        ev_badge = ("📎" if _ev.get("file") else "") + ("📝" if _ev.get("text") else "")

        rows.append({
            "ID": item_id,
            "Section": section_name,
            "Symptôme": tr_symptom(item),
            "Sévérité": severity,
            "Conformité": tr_conformite(conformite_actuelle),
            "Evidence": ev_badge or "—",
            "IA": ai_status,
            "Note": _ev.get("text", ""),
        })

    df_inspection = pd.DataFrame(rows)

    n_total = len(df_inspection)
    n_traites = sum(
        1 for v in df_inspection["Conformité"]
        if conformite_from_display(v) != CONFORMITE_A_INSPECTER
    )

    col1, col2 = st.columns([1, 3])
    with col1:
        st.metric(t("step1.progress"), f"{n_traites} / {n_total}")
    with col2:
        if n_total > 0:
            st.progress(n_traites / n_total)

    st.markdown("---")

    only_todo = st.checkbox(t("step1.only_todo"), key="insp_only_todo")
    df_view = df_inspection
    if only_todo:
        df_view = df_inspection[
            df_inspection["Conformité"].map(conformite_from_display) == CONFORMITE_A_INSPECTER
        ]

    edited_df = st.data_editor(
        df_view,
        use_container_width=True,
        hide_index=True,
        height=600,
        column_config={
            "ID": st.column_config.TextColumn(t("step1.col_id"), disabled=True, width="small"),
            "Section": st.column_config.TextColumn(t("step1.col_section"), disabled=True, width="medium"),
            "Symptôme": st.column_config.TextColumn(t("step1.col_symptom"), disabled=True, width="large"),
            "Sévérité": st.column_config.NumberColumn(t("step1.col_severity"), disabled=True, width="small"),
            "Conformité": st.column_config.SelectboxColumn(
                t("step1.col_conformity"),
                options=conformite_display_options(),
                required=True,
                width="small",
            ),
            "Evidence": st.column_config.TextColumn(
                t("step1.col_evidence"), disabled=True, width="small",
            ),
            "IA": st.column_config.TextColumn(
                t("step1.col_ai"), disabled=True, width="small",
                help=t("step1.col_ai_help"),
            ),
            "Note": st.column_config.TextColumn(t("step1.col_note"), width="medium"),
        },
        key="inspection_table_editor_" + _selection_key(st.session_state.get('insp_sections')) + ("_todo" if only_todo else ""),
    )

    st.markdown("---")
    col1, col2 = st.columns([1, 3])

    with col1:
        if st.button(
            t("step1.save_conformities"),
            type="primary",
            key="save_conformites_btn",
        ):
            n_saved = 0
            n_c = n_pc = n_nc = 0
            _items_by_id = {it[0]: it for it in all_items}
            orig_notes = dict(zip(df_inspection["ID"], df_inspection["Note"]))

            for _, row in edited_df.iterrows():
                item_id = row["ID"]
                conformite = conformite_from_display(row["Conformité"])

                if conformite == CONFORMITE_A_INSPECTER:
                    continue

                db.save_conformity_only(inspection_id, site_name, item_id, conformite)

                new_note = (row.get("Note") or "").strip()
                if new_note and new_note != (orig_notes.get(item_id) or "").strip():
                    db.save_inspection_data(
                        inspection_id, site_name, item_id,
                        {"conformite": conformite, "evidence_text": new_note},
                    )

                if conformite == "C":
                    n_c += 1
                elif conformite == "NC":
                    n_nc += 1
                else:
                    n_pc += 1

                item_row = _items_by_id.get(item_id) or db.get_item(item_id)
                if item_row:
                    st.session_state.inspection_buffer[item_id] = {
                        "conformite": conformite,
                        "severity": item_row[4],
                        "freq": item_row[5],
                        "method": item_row[6],
                        "t_human": bool(item_row[9]),
                        "t_env": bool(item_row[10]),
                        "t_asset": bool(item_row[11]),
                    }

                n_saved += 1

            st.success(t(
                "step1.save_success",
                n=n_saved, n_c=n_c, n_pc=n_pc, n_nc=n_nc,
            ))
            st.rerun()

    with col2:
        st.caption(t("step1.save_hint"))

    if n_traites > 0:
        st.markdown("---")
        st.subheader(t("step1.recap"))

        def _count_conf(code):
            return sum(
                1 for v in df_inspection["Conformité"]
                if conformite_from_display(v) == code
            )

        col1, col2, col3, col4 = st.columns(4)
        col1.metric(t("step1.total_processed"), n_traites)
        col2.metric(tr_conformite("C"), _count_conf("C"))
        col3.metric(tr_conformite("PC"), _count_conf("PC"))
        col4.metric(tr_conformite("NC"), _count_conf("NC"))

    # ---- EVIDENCE ----
    st.markdown("---")
    st.subheader(t("step1.evidence_title"))
    st.caption(t("step1.evidence_caption"))

    # all_items = symptômes des sections choisies (voir _pick_domain_and_sections plus haut)
    conformites_actuelles = db.get_conformity_by_inspection(inspection_id)

    existing_inspection_data = db.get_inspection_data(inspection_id)
    evidence_existing = {}
    for e in existing_inspection_data:
        item_id_e = e[3]
        has_text = bool(e[25]) if len(e) > 25 else False
        has_file = bool(e[26]) if len(e) > 26 else False
        evidence_existing[item_id_e] = {
            "text": e[25] if len(e) > 25 else "",
            "file": e[26] if len(e) > 26 else None,
            "has_text": has_text,
            "has_file": has_file,
        }

    col1, col2 = st.columns([2, 1])

    with col1:
        search_ev = st.text_input(
            t("step1.evidence_search"),
            key="evidence_search",
            placeholder=t("step1.evidence_search_placeholder"),
        )

    with col2:
        FILTRES_EV = ["ALL", "ONLY_C", "ONLY_PC", "ONLY_NC", "NO_EV", "WITH_EV"]
        _FILTRE_EV_KEYS = {
            "ALL": "step1.filter_all",
            "ONLY_C": "step1.filter_only_c",
            "ONLY_PC": "step1.filter_only_pc",
            "ONLY_NC": "step1.filter_only_nc",
            "NO_EV": "step1.filter_no_evidence",
            "WITH_EV": "step1.filter_with_evidence",
        }
        filtre_ev = st.selectbox(
            t("step1.evidence_filter"),
            FILTRES_EV,
            format_func=lambda x: t(_FILTRE_EV_KEYS[x]),
            key="evidence_filter",
        )

    items_filtres = []
    for item in all_items:
        iid = item[0]
        symptome_i = symptom_search_text(item)
        evidence_requis = item[8] if len(item) > 8 else ""
        conf_i = conformites_actuelles.get(iid, CONFORMITE_A_INSPECTER)

        if search_ev:
            search_lower = search_ev.lower()
            if (search_lower not in iid.lower() and
                search_lower not in symptome_i.lower() and
                search_lower not in evidence_requis.lower()):
                continue

        ev = evidence_existing.get(iid, {})
        has_ev = ev.get("has_text", False) or ev.get("has_file", False)

        if filtre_ev == "ONLY_C" and conf_i != "C":
            continue
        elif filtre_ev == "ONLY_PC" and conf_i != "PC":
            continue
        elif filtre_ev == "ONLY_NC" and conf_i != "NC":
            continue
        elif filtre_ev == "NO_EV" and has_ev:
            continue
        elif filtre_ev == "WITH_EV" and not has_ev:
            continue

        items_filtres.append(item)

    st.markdown(f"**{t('step1.n_symptoms_shown', n=len(items_filtres))}**")

    evidence_dir = Path("inspection_evidence") / inspection_id
    evidence_dir.mkdir(parents=True, exist_ok=True)

    for item in items_filtres:
        iid = item[0]
        symptome_i = tr_symptom(item)
        evidence_requis = item[8] if len(item) > 8 else ""
        conf_i = conformites_actuelles.get(iid, CONFORMITE_A_INSPECTER)

        emoji = CONFORMITE_EMOJI.get(conf_i, "⚪")

        ev = evidence_existing.get(iid, {})
        has_text = ev.get("has_text", False)
        has_file = ev.get("has_file", False)

        badge = ""
        if has_file:
            badge += " 📎"
        if has_text:
            badge += " 📝"

        _sec_tag = ""
        if len(st.session_state.get("insp_sections") or []) > 1:
            _sid_i = item_section_id(item)
            _sec_tag = f"[{_sid_i}] " if _sid_i is not None else ""
        titre = f"{emoji} {_sec_tag}{iid} — {symptome_i[:70]}{badge}"

        with st.expander(titre, expanded=False):
            st.markdown(f"**{t('step1.col_equipment')} :** {tr_equipment(item)}")
            st.markdown(f"**{t('step1.evidence_required')} :** {evidence_requis}")
            st.markdown("---")

            ai_saved = db.get_ai_verification(inspection_id, iid)
            if ai_saved:
                verdict_saved = ai_saved["verdict"]
                score_saved = ai_saved["score"]
                emoji_saved = AI_EMOJI.get(verdict_saved, "❓")

                st.markdown(t(
                    "step1.ai_verdict_line",
                    emoji=emoji_saved,
                    verdict=t(f"step1.ai_verdict_{verdict_saved}"),
                    score=score_saved,
                ))
                if ai_saved.get("commentaire"):
                    st.caption(f"💬 {ai_saved['commentaire']}")

                if ai_saved.get("suggestions"):
                    with st.expander(t("step1.ai_suggestions_saved")):
                        for s in ai_saved["suggestions"]:
                            st.markdown(f"- {s}")

                st.markdown("---")

            note_val = ev.get("text", "")
            note = st.text_area(
                t("step1.evidence_note"),
                value=note_val,
                key=f"ev_text_{iid}",
                height=80,
                placeholder=t("step1.evidence_note_placeholder"),
            )

            uploaded_file = st.file_uploader(
                t("step1.evidence_file"),
                type=["png", "jpg", "jpeg", "pdf", "docx", "xlsx", "txt"],
                key=f"ev_file_{iid}",
            )

            if has_file:
                st.caption(t("step1.evidence_current_file", file=ev.get("file")))

            if st.button(t("step1.evidence_save", id=iid), key=f"save_ev_{iid}"):
                file_path = ev.get("file")

                if uploaded_file is not None:
                    safe_name = uploaded_file.name.replace(" ", "_")
                    file_path = evidence_dir / f"{iid}_{safe_name}"
                    with open(file_path, "wb") as f:
                        f.write(uploaded_file.getbuffer())

                inspection_data = {
                    "conformite": conf_i if conf_i != CONFORMITE_A_INSPECTER else None,
                    "evidence_text": note,
                    "evidence_file": str(file_path) if file_path else None,
                }

                for e in existing_inspection_data:
                    if e[3] == iid:
                        inspection_data.update({
                            "n_personnes": e[5],
                            "heures_expo": e[6],
                            "p_fat": e[7],
                            "substance": e[8],
                            "taille_fuite_kg": e[9],
                            "duree_fuite_s": e[10],
                            "vitesse_vent_ms": e[11],
                            "direction_vent_deg": e[12],
                            "temperature_c": e[13],
                            "pression_hpa": e[14],
                            "humidite_pct": e[15],
                            "latitude": e[16],
                            "longitude": e[17],
                            "source_meteo": e[18],
                            "cout_action": e[19],
                            "cout_perte": e[20],
                            "duree_arret_j": e[21],
                            "cout_indispo_jour": e[22],
                            "currency": e[23],
                            "industry": e[24],
                        })
                        break

                db.save_inspection_data(inspection_id, site_name, iid, inspection_data)
                st.success(t("step1.evidence_saved", id=iid))
                if uploaded_file is not None:
                    st.caption(t("step1.evidence_file_saved", file=file_path))

                if not ai_service.is_configured():
                    st.info(t("step1.ai_not_configured"))
                else:
                    with st.spinner(t("step1.ai_verifying")):
                        ai_result = ai_service.verifier_evidence(
                            symptome_id=iid,
                            symptome_description=symptome_i,
                            evidence_requis=evidence_requis,
                            evidence_text=note,
                            image_path=str(file_path) if file_path and uploaded_file else None,
                            conformite=conf_i if conf_i != CONFORMITE_A_INSPECTER else "NC",
                        )

                    if ai_result.get("erreur"):
                        st.warning(t("step1.ai_error", error=ai_result["erreur"]))
                    else:
                        db.save_ai_verification(inspection_id, iid, ai_result)

                        verdict = ai_result["verdict"]
                        score = ai_result["score"]
                        emoji = AI_EMOJI.get(verdict, "❓")
                        line = t(
                            "step1.ai_verdict_line",
                            emoji=emoji,
                            verdict=t(f"step1.ai_verdict_{verdict}"),
                            score=score,
                        )

                        if verdict == "conforme":
                            st.success(line)
                        elif verdict == "partiel":
                            st.warning(line)
                        else:
                            st.error(line)

                        if ai_result.get("commentaire"):
                            st.markdown(f"**{t('step1.ai_comment_label')} :** {ai_result['commentaire']}")

                        if ai_result.get("suggestions"):
                            with st.expander(t("step1.ai_suggestions")):
                                for s in ai_result["suggestions"]:
                                    st.markdown(f"- {s}")


# ============================================================
# ÉTAPE 2 — ANALYSE QRA (Monte Carlo + Dispersion)
# ============================================================
def calculer_score_matrice(sev_n, freq_moyenne, df_prob):
    prob_level, _, _ = freq_to_prob_level(freq_moyenne, df_prob)
    return float(sev_n) * int(prob_level), int(prob_level)


def render_step_2_qra():
    """Étape 2 : Analyse QRA (Monte Carlo batch → Top 10 → Formulaires)."""

    st.header(t("step2.title"))

    site_name = st.session_state.get("active_site")
    if not site_name:
        st.error(t("step2.no_active_matrix"))
        return

    try:
        matrix_config = RiskMatrixConfig(site_name=site_name)
        _, df_prob, _ = load_matrix_dfs(site_name)
    except Exception as e:
        st.error(t("common.matrix_error", error=e))
        return

    buffer = st.session_state.inspection_buffer
    nc_pc_items = {
        iid: data for iid, data in buffer.items()
        if data["conformite"] in ("NC", "PC")
    }

    if not nc_pc_items:
        st.info(t("step2.no_nc_pc"))
        return

    st.caption(t("step2.nc_pc_count", n=len(nc_pc_items), site=site_name))
    st.markdown("---")

    # PHASE A
    st.subheader(t("step2.phase_a_title"))
    st.caption(t("step2.phase_a_caption"))

    phase_a_done = st.session_state.get("phase_a_done", False)
    col1, col2 = st.columns([1, 3])

    with col1:
        if st.button(t("step2.run_quick"), type="primary", key="run_phase_a_btn"):
            progress = st.progress(0)
            status = st.empty()
            results_a = []

            for idx, (iid, data) in enumerate(nc_pc_items.items()):
                status.text(t("step2.analyzing", i=idx + 1, total=len(nc_pc_items), id=iid))

                item_row = db.get_item(iid)
                if not item_row:
                    continue

                sev_n = item_row[4]

                try:
                    freq_ref = data.get("freq", 0.001)
                    method_n = data.get("method", "Estimated")
                    factor = {"Direct": 2.0, "Calcule": 3.0, "Estimated": 5.0}.get(method_n, 3.0)

                    echant = simuler_frequence_triangulaire(
                        freq_ref / factor, freq_ref, freq_ref * factor,
                        n_iterations=N_ITER_APERCU,
                    )
                    stats = resumer_distribution(echant)
                    score, prob_level = calculer_score_matrice(
                        sev_n, stats["freq_moyenne"], df_prob
                    )

                    results_a.append({
                        "item_id": iid,
                        "conformite": data["conformite"],
                        "severite": sev_n,
                        "freq_moyenne": stats["freq_moyenne"],
                        "freq_p50": stats["freq_p50"],
                        "freq_p95": stats["freq_p95"],
                        "score": score,
                        "prob_level": prob_level,
                        "buffer_data": data,
                    })
                except Exception as e:
                    st.warning(t("step2.item_error", id=iid, error=e))

                progress.progress((idx + 1) / len(nc_pc_items))

            progress.empty()
            status.empty()

            results_a.sort(
                key=lambda r: (r["score"], r["freq_moyenne"] * float(r["severite"])),
                reverse=True,
            )
            top_10 = results_a[:10]

            st.session_state["phase_a_results"] = results_a
            st.session_state["top_10"] = top_10
            st.session_state["phase_a_done"] = True

            st.success(t("step2.quick_done", n=len(results_a)))
            st.rerun()

    with col2:
        if phase_a_done:
            st.info(t("step2.already_done", n=len(st.session_state.get("top_10", []))))
        else:
            st.caption(t("step2.quick_hint"))

    if not phase_a_done:
        st.markdown("---")
        st.info(t("step2.run_phase_a_first"))
        return

    # PHASE B
    st.markdown("---")
    st.subheader(t("step2.phase_b_title"))

    top_10 = st.session_state.get("top_10", [])
    if not top_10:
        st.warning(t("step2.no_top10"))
        return

    df_top10 = pd.DataFrame([
        {
            t("step2.col_rank"): i + 1,
            t("step2.col_id"): r["item_id"],
            t("step2.col_conformity"): tr_conformite(r["conformite"]),
            t("step2.col_severity"): r["severite"],
            t("step2.col_freq_mean"): f"{r['freq_moyenne']:.6e}",
            t("step2.col_freq_p95"): f"{r['freq_p95']:.6e}",
            t("step2.col_prob_level"): r.get("prob_level"),
            t("step2.col_score"): f"{r['score']:.0f}",
        }
        for i, r in enumerate(top_10)
    ])

    st.dataframe(df_top10, use_container_width=True, hide_index=True)
    st.info(t("step2.forms_below"))

    # PHASE C
    st.markdown("---")
    st.subheader(t("step2.phase_c_title"))
    st.caption(t("step2.phase_c_caption"))

    for rank, top_item in enumerate(top_10, start=1):
        iid = top_item["item_id"]
        data = top_item["buffer_data"]
        conformite = top_item["conformite"]

        item_row = db.get_item(iid)
        if not item_row:
            continue

        section_n = tr_section(item_row)
        equip_n = tr_equipment(item_row)
        sympt_n = tr_symptom(item_row)
        sev_n = item_row[4]
        freq_n = item_row[5]
        method_n = item_row[6]
        t_h = item_row[9]
        t_e = item_row[10]
        t_a = item_row[11]

        t_human = bool(t_h)
        t_env = bool(t_e)
        t_asset = bool(t_a)

        emoji = CONFORMITE_EMOJI.get(conformite, "⚠️")

        with st.expander(
            t("step2.form_expander",
              rank=rank, id=iid, emoji=emoji,
              conformite=tr_conformite(conformite),
              score=f"{top_item['score']:.4e}"),
            expanded=(rank == 1),
        ):
            st.markdown(f"""
**{t('step2.ctx_section')}** : {section_n}  
**{t('step2.ctx_equipment')}** : {equip_n}  
**{t('step2.ctx_symptom')}** : {sympt_n[:200]}  
**{t('step2.ctx_severity')}** : {sev_n} | **{t('step2.ctx_freq_cat')}** : {freq_n} | **{t('step2.ctx_method')}** : {tr_method_of(item_row)}
""")
            st.markdown("---")

            # MC
            st.markdown(f"##### {t('step2.mc_final_title')}")
            _freq_ref = data.get("freq", 0.001)
            _factor = {"Direct": 2.0, "Calcule": 3.0, "Estimated": 5.0}.get(method_n, 3.0)
            _freq_min = _freq_ref / _factor
            _freq_mode = _freq_ref
            _freq_max = _freq_ref * _factor

            st.caption(t("step2.mc_bounds",
                         a=f"{_freq_min:.6e}",
                         m=f"{_freq_mode:.6e}",
                         b=f"{_freq_max:.6e}"))

            override_freq = st.checkbox(t("step2.mc_override"), value=False, key=f"override_{iid}")

            if override_freq:
                col1, col2, col3 = st.columns(3)
                fmin = col1.number_input(t("step2.mc_a_min"), value=_freq_min, format="%.8f", key=f"fmin_{iid}")
                fmode = col2.number_input(t("step2.mc_m_mode"), value=_freq_mode, format="%.8f", key=f"fmode_{iid}")
                fmax = col3.number_input(t("step2.mc_b_max"), value=_freq_max, format="%.8f", key=f"fmax_{iid}")
            else:
                fmin, fmode, fmax = _freq_min, _freq_mode, _freq_max

            st.markdown("")

            # HUMAIN
            n_personnes = heures_expo = p_fat = None
            if t_human:
                st.markdown(f"##### {t('step2.human_title')}")
                col1, col2, col3 = st.columns(3)
                n_personnes = col1.number_input(t("step2.human_n_people"), min_value=0, value=5, key=f"npers_{iid}")
                heures_expo = col2.number_input(t("step2.human_expo_hours"), min_value=0.0, value=2000.0, key=f"hexpo_{iid}")
                p_fat = col3.number_input(t("step2.human_p_fat"), min_value=0.0, max_value=1.0, value=0.1, step=0.01, key=f"pfat_{iid}")
                st.markdown("")

            # ENV — WIZARD
            substance_env = None
            taille_fuite = duree_fuite = None
            vitesse_vent = direction_vent = temperature = None
            latitude = longitude = None
            source_meteo = "manuelle"
            use_weather = False

            if t_env:
                st.markdown(f"##### {t('step2.env_title')}")

                # Appel au wizard complet
                _wizard_result = render_dispersion_wizard(
                    iid=iid,
                    substance_options=substance_options,
                )
                if _wizard_result:
                    substance_env = _wizard_result.get("substance")
                    taille_fuite = _wizard_result.get("taille_fuite")
                    duree_fuite = _wizard_result.get("duree_fuite")
                    vitesse_vent = _wizard_result.get("vitesse_vent")
                    direction_vent = _wizard_result.get("direction_vent")
                    temperature = _wizard_result.get("temperature")
                    latitude = _wizard_result.get("latitude")
                    longitude = _wizard_result.get("longitude")
                    source_meteo = _wizard_result.get("source_meteo", "manuelle")
                    use_weather = _wizard_result.get("use_weather", False)

                st.markdown("")

                # Bloc GHS
                if substance_env and substance_env != "-- Sélectionner --":
                    with st.expander("⚠️ Classification GHS", expanded=False):
                        render_ghs_block(substance_env)

                st.markdown("")

            # ASSET
            cout_action = cout_perte = duree_arret = cout_indispo_jour = None
            currency = DEFAULT_CURRENCY
            industry = "generic"

            if t_asset:
                st.markdown(f"##### {t('step2.asset_title')}")

                col1, col2 = st.columns(2)
                currency = col1.selectbox(
                    t("step2.asset_currency"),
                    list(CURRENCIES.keys()),
                    index=list(CURRENCIES.keys()).index(DEFAULT_CURRENCY),
                    format_func=lambda c: f"{c} — {CURRENCIES[c]['label']}",
                    key=f"curr_{iid}",
                )

                symbol = CURRENCIES[currency]["symbol"]

                col1, col2 = st.columns(2)
                cout_action = col1.number_input(
                    t("step2.asset_action_cost", symbol=symbol),
                    min_value=0.0, value=15000.0,
                    key=f"cact_{iid}",
                )
                cout_perte = col2.number_input(
                    t("step2.asset_loss_cost", symbol=symbol),
                    min_value=0.0, value=200000.0,
                    key=f"cperte_{iid}",
                )

                col1, col2 = st.columns(2)
                duree_arret = col1.number_input(
                    t("step2.asset_downtime_days"), min_value=0.0, value=10.0,
                    key=f"darret_{iid}",
                )

                default_indispo = get_default_indisponibilite(industry)
                if currency != "EUR":
                    default_indispo = convert_from_eur(default_indispo, currency)

                cout_indispo_jour = col2.number_input(
                    t("step2.asset_downtime_cost", symbol=symbol),
                    min_value=0.0, value=float(default_indispo),
                    key=f"cindispo_{iid}",
                )
                st.markdown("")

            # SAVE FORM
            st.markdown("---")
            if st.button(t("step2.save_form", id=iid), type="primary", key=f"save_form_{iid}"):
                st.session_state.inspection_buffer[iid].update({
                    "freq_min": fmin, "freq_mode": fmode, "freq_max": fmax,
                    "n_personnes": n_personnes, "heures_expo": heures_expo, "p_fat": p_fat,
                    "substance_env": substance_env,
                    "taille_fuite": taille_fuite, "duree_fuite": duree_fuite,
                    "vitesse_vent": vitesse_vent, "direction_vent": direction_vent,
                    "temperature": temperature,
                    "latitude": latitude, "longitude": longitude,
                    "source_meteo": source_meteo,
                    "cout_action": cout_action, "cout_perte": cout_perte,
                    "duree_arret": duree_arret, "cout_indispo_jour": cout_indispo_jour,
                    "currency": currency, "industry": industry,
                })
                st.success(t("step2.form_saved", id=iid))

    # FINAL
    st.markdown("---")
    st.subheader(t("step2.final_title"))
    st.caption(t("step2.final_caption"))

    if st.button(t("step2.final_button"), type="primary", key="run_final_analysis_btn"):
        progress = st.progress(0)
        status = st.empty()

        for idx, top_item in enumerate(top_10):
            iid = top_item["item_id"]
            data = st.session_state.inspection_buffer.get(iid)
            if not data:
                continue

            status.text(t("step2.final_analyzing", i=idx + 1, total=len(top_10), id=iid))

            item_row = db.get_item(iid)
            sev_n = item_row[4]
            t_h = item_row[9]
            t_e = item_row[10]
            t_a = item_row[11]

            freq_min_final = data.get("freq_min", data["freq"])
            freq_mode_final = data.get("freq_mode", data["freq"])
            freq_max_final = data.get("freq_max", data["freq"] * 3)

            echant = simuler_frequence_triangulaire(
                freq_min_final, freq_mode_final, freq_max_final,
                n_iterations=N_ITER_FINAL,
            )
            stats = resumer_distribution(echant)
            position = matrix_config.positionner(sev_n, stats["freq_moyenne"])
            score, _ = calculer_score_matrice(sev_n, stats["freq_moyenne"], df_prob)

            inspection_data = {
                "conformite": data["conformite"],
                "n_personnes": data.get("n_personnes"),
                "heures_expo": data.get("heures_expo"),
                "p_fat": data.get("p_fat"),
                "substance": data.get("substance_env"),
                "taille_fuite_kg": data.get("taille_fuite"),
                "duree_fuite_s": data.get("duree_fuite"),
                "vitesse_vent_ms": data.get("vitesse_vent"),
                "direction_vent_deg": data.get("direction_vent"),
                "temperature_c": data.get("temperature"),
                "latitude": data.get("latitude"),
                "longitude": data.get("longitude"),
                "source_meteo": data.get("source_meteo"),
                "cout_action": data.get("cout_action"),
                "cout_perte": data.get("cout_perte"),
                "duree_arret_j": data.get("duree_arret"),
                "cout_indispo_jour": data.get("cout_indispo_jour"),
                "currency": data.get("currency"),
                "industry": data.get("industry"),
            }
            db.save_inspection_data(inspection_id, site_name, iid, inspection_data)

            db.save_mc_result(inspection_id, iid, {
                "n_iterations": N_ITER_FINAL,
                "freq_min": freq_min_final,
                "freq_mode": freq_mode_final,
                "freq_max": freq_max_final,
                **stats,
                "severite": sev_n,
                "score_risque": score,
                "niveau_severite_matrice": position["label_severite"],
                "niveau_probabilite_matrice": position["label_probabilite"],
                "niveau_risque_matrice": position["niveau_risque"],
                "couleur_matrice": position["couleur"],
            })

            if t_h and data.get("n_personnes"):
                try:
                    hr = evaluer_risque_humain(
                        symptom_id=iid,
                        frequence_par_an=stats["freq_moyenne"],
                        nb_personnes_exposees=data["n_personnes"],
                        duree_exposition_h_an=data["heures_expo"],
                        proba_deces=data["p_fat"],
                    )
                    db.save_human_risk(inspection_id, iid, hr)
                except Exception:
                    pass

            if t_a and data.get("cout_action") is not None:
                try:
                    cb = evaluer_cout_benefice(
                        symptom_id=iid,
                        frequence_par_an=stats["freq_moyenne"],
                        cout_action_corrective=data["cout_action"],
                        cout_perte_potentielle=data.get("cout_perte", 0.0),
                        duree_indisponibilite_j=data.get("duree_arret", 0.0),
                        cout_indisponibilite_par_jour=data.get("cout_indispo_jour", 0.0),
                        currency=data.get("currency", DEFAULT_CURRENCY),
                        industry=data.get("industry", "generic"),
                    )
                    db.save_cost_benefit(inspection_id, iid, cb)
                except Exception as e:
                    st.error(t("step2.cost_benefit_failed", id=iid, error=e))

            if t_e and data.get("substance_env") and data["substance_env"] != "-- Sélectionner --":
                try:
                    disp = evaluer_dispersion(
                        symptom_id=iid,
                        substance=data["substance_env"],
                        taille_fuite_kg=data["taille_fuite"],
                        duree_fuite_s=data["duree_fuite"],
                        vitesse_vent_m_s=data["vitesse_vent"],
                        direction_vent_deg=data["direction_vent"],
                        temperature_c=data["temperature"],
                        source_meteo=data.get("source_meteo", "manuelle"),
                    )
                    db.save_dispersion(inspection_id, iid, disp)
                except Exception:
                    pass

            progress.progress((idx + 1) / len(top_10))

        progress.empty()
        status.empty()
        st.success(t("step2.final_done", n=len(top_10)))
        st.session_state.batch_done = True
        st.info(t("step2.goto_step3"))


# ============================================================
# ÉTAPE 3 — RÉSULTATS + EMAIL + EXCEL + PDF
# ============================================================
def render_step_3_results():
    """Étape 3 : Résultats + classification + email + export Excel/PDF."""

    st.header(t("step3.title"))

    if not st.session_state.batch_done:
        st.info(t("step3.batch_not_run"))
        return

    mc_results = db.get_mc_results(inspection_id)
    if not mc_results:
        st.warning(t("step3.no_results"))
        return

    site_name = st.session_state.get("active_site")
    if not site_name:
        st.error(t("step3.no_active_matrix"))
        return

    try:
        df_sev, df_prob, df_seuils = load_matrix_dfs(site_name)
    except Exception as e:
        st.error(t("step3.matrix_load_error", error=e))
        return

    st.caption(t("step3.active_matrix", site=site_name))

    df = pd.DataFrame(mc_results, columns=[
        "id", "inspection_id", "item_id", "n_iterations",
        "freq_min", "freq_mode", "freq_max",
        "freq_moyenne", "freq_p05", "freq_p50", "freq_p95",
        "severite", "score_risque",
        "niveau_sev", "niveau_prob", "niveau_risque", "couleur", "created_at",
    ])

    df_classified = classify_results_df(df, df_sev, df_prob, df_seuils)
    df_classified = df_classified.sort_values("risk_score", ascending=False)
    df_classified["rank"] = range(1, len(df_classified) + 1)

    st.markdown("---")
    st.subheader(t("step3.overview"))
    render_metric_cards(df_classified, df_seuils)
    st.markdown("---")

    # ---- EMAIL BUTTON ----
    email_status = email_service.get_status()
    if (email_status["configured"] and st.session_state.email_recipient
            and auth.has_permission("send_email")):
        st.markdown("### 📧 Envoyer le rapport")
        col1, col2 = st.columns([2, 1])

        with col1:
            st.caption(f"Destinataire : **{st.session_state.email_recipient}**")

        with col2:
            if st.button("📧 Envoyer maintenant", key="send_report_from_step3",
                         type="primary", use_container_width=True):
                with st.spinner("Préparation et envoi..."):
                    conformites = db.get_conformity_by_inspection(inspection_id)
                    n_c = sum(1 for v in conformites.values() if v == "C")
                    n_pc = sum(1 for v in conformites.values() if v == "PC")
                    n_nc = sum(1 for v in conformites.values() if v == "NC")

                    top5_html = "<ul>"
                    for _, row in df_classified.head(5).iterrows():
                        top5_html += f"<li><b>{row['item_id']}</b> — Score : {row['risk_score']:.2f} — Classe : {row['risk_class']}</li>"
                    top5_html += "</ul>"

                    report_body = f"""
                    <h2>📊 Résultats de classification</h2>
                    <p><b>Site :</b> {site_name} | <b>Référence :</b> {inspection_id}</p>
                    <h3>🏆 Top 5 des risques</h3>
                    {top5_html}
                    <h3>📋 Classification complète</h3>
                    <p>{len(df_classified)} symptôme(s) analysé(s) — voir rapport complet dans HSE PYXIS.</p>
                    """

                    r = email_service.send_report(
                        to=st.session_state.email_recipient,
                        site_name=site_name,
                        inspection_id=inspection_id,
                        html_report=report_body,
                        summary={
                            "total": len(conformites),
                            "nc": n_nc,
                            "pc": n_pc,
                            "c": n_c,
                        },
                    )

                if r["ok"]:
                    st.success(f"✅ Rapport envoyé via **{r['provider']}**")
                else:
                    st.error(f"❌ Échec : {r['error']}")

        st.markdown("---")

    # ---- EXPORT EXCEL + PDF ----
    st.markdown("### 💾 Télécharger les résultats")

    hr_rows = db.get_human_risk(inspection_id)
    df_hr = pd.DataFrame(hr_rows, columns=[
        "id", "inspection_id", "item_id", "frequence_par_an",
        "nb_personnes_exposees", "duree_exposition_h_an",
        "proba_deces", "proba_deces_est_defaut",
        "pll", "far", "lira", "created_at",
    ])

    cb_rows = db.get_cost_benefit(inspection_id)
    df_cb = pd.DataFrame(cb_rows, columns=[
        "id", "inspection_id", "item_id", "currency", "industry",
        "frequence_par_an", "cout_action_corrective",
        "cout_perte_potentielle", "duree_indisponibilite_j",
        "cout_indisponibilite_par_jour", "perte_annuelle_attendue",
        "ratio_benefice_cout", "recommandation", "created_at",
    ])

    disp_rows = db.get_dispersion(inspection_id)
    df_disp = pd.DataFrame(disp_rows, columns=[
        "id", "inspection_id", "item_id", "substance", "taille_fuite_kg",
        "duree_fuite_s", "vitesse_vent_ms", "direction_vent_deg",
        "temperature_c", "pression_hpa", "humidite_pct",
        "latitude", "longitude", "classe_stabilite",
        "distance_impact_m", "source_meteo", "created_at",
    ])

    col_xl, col_pdf = st.columns(2)

    # Excel
    with col_xl:
        try:
            excel_bytes = export_to_excel(
                df_classified=df_classified,
                df_hr=df_hr,
                df_cb=df_cb,
                df_disp=df_disp,
            )
            st.download_button(
                label=t("exports.excel_button"),
                data=excel_bytes,
                file_name=f"inspection_{inspection_id}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key=f"dl_excel_{inspection_id}",
            )
        except Exception as e:
            st.warning(t("exports.excel_error", error=e))

    # PDF
    with col_pdf:
        try:
            pdf_bytes = export_to_pdf(
                inspection_id=str(inspection_id),
                site_name=site_name,
                inspector_name=current_user.get("full_name", current_user.get("username", "")),
                results_df=df_classified,
                summary={
                    "NC": sum(1 for v in db.get_conformity_by_inspection(inspection_id).values() if v == "NC"),
                    "PC": sum(1 for v in db.get_conformity_by_inspection(inspection_id).values() if v == "PC"),
                    "C": sum(1 for v in db.get_conformity_by_inspection(inspection_id).values() if v == "C"),
                    "Top risque": df_classified.iloc[0]["item_id"] if not df_classified.empty else "—",
                },
                language=st.session_state.get("lang", "fr"),
            )
            if pdf_bytes:
                st.download_button(
                    label=t("exports.pdf_button"),
                    data=pdf_bytes,
                    file_name=f"rapport_{inspection_id}.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                    key=f"dl_pdf_{inspection_id}",
                )
                if auth.has_permission("edit_inspection"):
                    auth.audit("EXPORT_PDF", "inspection", str(inspection_id), None)
            else:
                st.caption(t("exports.pdf_unavailable"))
        except Exception as e:
            st.warning(t("exports.pdf_error", error=e))

    st.markdown("---")

    # ---- RISQUE HUMAIN ----
    if not df_hr.empty:
        df_hr = df_hr.merge(
            df_classified[["item_id", "severite", "risk_class", "risk_color"]],
            on="item_id", how="left",
        )

    if df_cb.empty:
        st.warning(t("step3.no_cost_benefit"))

    tab_labels = [
        t("step3.tab_top10"),
        t("step3.tab_matrix"),
        t("step3.tab_all"),
        t("step3.tab_plots"),
        t("step3.tab_human"),
    ]
    if not df_cb.empty:
        tab_labels.append(t("step3.tab_cost_benefit"))
    if not df_disp.empty:
        tab_labels.append(t("step3.tab_dispersion"))

    _tabs = st.tabs(tab_labels)
    tab_top, tab_matrix, tab_all, tab_plots, tab_human = _tabs[:5]
    _next = 5
    tab_cb = None
    tab_disp = None
    if not df_cb.empty:
        tab_cb = _tabs[_next]
        _next += 1
    if not df_disp.empty:
        tab_disp = _tabs[_next]

    with tab_top:
        st.subheader(t("step3.top10_title"))
        top10 = df_classified.head(10)[
            ["rank", "item_id", "freq_moyenne", "freq_p95",
             "severite", "prob_label", "risk_score", "risk_class", "risk_color"]
        ]
        top10_display = top10.drop(columns=["risk_color"]).copy()
        top10_display["prob_label"] = top10_display["prob_label"].apply(tr_matrix_label)
        top10_display["risk_class"] = top10_display["risk_class"].apply(tr_matrix_label)
        top10_display = top10_display.rename(columns={
            "rank": t("step3.col_rank"),
            "item_id": t("step3.col_symptom"),
            "freq_moyenne": t("step3.col_freq_mean"),
            "freq_p95": t("step3.col_freq_p95"),
            "severite": t("step3.col_severity"),
            "prob_label": t("step3.col_probability"),
            "risk_score": t("step3.col_risk_score"),
            "risk_class": t("step3.col_risk_class"),
        })
        st.dataframe(top10_display, use_container_width=True, hide_index=True)

    with tab_matrix:
        st.subheader(t("step3.matrix_title"))
        st.caption(t("step3.matrix_caption", site=site_name))
        render_heatmap(df_classified, df_sev, df_prob, df_seuils)

    with tab_all:
        st.subheader(t("step3.all_title"))
        available_classes = df_seuils["risk_label"].tolist()
        filter_class = st.multiselect(
            t("step3.filter_by_class"),
            available_classes,
            default=available_classes,
            format_func=tr_matrix_label,
            key="filter_class_results",
        )
        df_display = df_classified[df_classified["risk_class"].isin(filter_class)]
        df_display_tr = df_display[[
            "rank", "item_id", "severite", "freq_moyenne",
            "prob_label", "risk_score", "risk_class",
        ]].copy()
        df_display_tr["prob_label"] = df_display_tr["prob_label"].apply(tr_matrix_label)
        df_display_tr["risk_class"] = df_display_tr["risk_class"].apply(tr_matrix_label)
        st.dataframe(
            df_display_tr.rename(columns={
                "rank": t("step3.col_rank"),
                "item_id": t("step3.col_symptom"),
                "severite": t("step3.col_severity"),
                "freq_moyenne": t("step3.col_freq_mean"),
                "prob_label": t("step3.col_probability"),
                "risk_score": t("step3.col_risk_score"),
                "risk_class": t("step3.col_risk_class"),
            }),
            use_container_width=True,
            hide_index=True,
        )
        csv = df_classified.to_csv(index=False).encode("utf-8")
        st.download_button(
            t("step3.download_csv"),
            csv,
            f"inspection_{inspection_id}_results.csv",
            "text/csv",
        )

    with tab_plots:
        st.subheader(t("step3.plots_title"))
        render_plots(df_classified, df_seuils, top_n=10)

    with tab_human:
        st.subheader(t("step3.human_title"))
        if df_hr.empty:
            st.info(t("step3.human_empty"))
        else:
            df_hr_disp = df_hr.copy()
            df_hr_disp["niveau_lira_code"] = df_hr_disp["lira"].apply(niveau_lira)
            df_hr_disp["niveau_lira"] = df_hr_disp["niveau_lira_code"].apply(tr_lira)
            st.caption(t(
                "step3.lira_thresholds",
                high=f"{SEUILS_LIRA_INDICATIFS['inacceptable']:.0e}",
                low=f"{SEUILS_LIRA_INDICATIFS['negligeable']:.0e}",
            ))
            st.dataframe(
                df_hr_disp[[
                    "item_id", "nb_personnes_exposees", "duree_exposition_h_an",
                    "proba_deces", "pll", "far", "lira", "niveau_lira",
                ]].rename(columns={
                    "item_id": t("step3.col_symptom"),
                    "nb_personnes_exposees": t("step3.col_n_people"),
                    "duree_exposition_h_an": t("step3.col_expo_hours"),
                    "proba_deces": t("step3.col_p_death"),
                    "pll": t("step3.col_pll"),
                    "far": t("step3.col_far"),
                    "lira": t("step3.col_lira"),
                    "niveau_lira": t("step3.col_lira_level"),
                }),
                use_container_width=True,
                hide_index=True,
            )

    if tab_cb is not None:
        with tab_cb:
            st.subheader(t("step3.cb_title"))
            st.caption(t("step3.cb_caption"))

    if tab_disp is not None:
        with tab_disp:
            st.subheader(t("step3.disp_title"))

    # ---- 📅 Inspection de suivi ----
    st.markdown("---")
    _confs = db.get_conformity_by_inspection(inspection_id)
    _nc_ids = [i for i, c in _confs.items() if c == "NC"]
    _pc_ids = [i for i, c in _confs.items() if c == "PC"]
    _sec_names = sorted({
        tr_section(it_) for it_ in (db.get_item(i) for i in _nc_ids + _pc_ids) if it_
    })
    if auth.has_permission("edit_schedule"):
        st.button(
            t("calendar.create_from_report"),
            key="create_insp_from_report",
            type="primary",
            on_click=_cb_plan_followup,
            args=(site_name, _sec_names,
                  t("calendar.followup_notes", id=inspection_id,
                    n_nc=len(_nc_ids), n_pc=len(_pc_ids))),
        )


# ============================================================
# ÉTAPE 4 — PLAN PDCA (IA)
# ============================================================
def _render_pdca_plan():
    """Étape 4 : PDCA (plan généré par l'IA)."""

    st.header(t("step5.title"))
    st.caption(t("step5.caption"))

    if not st.session_state.batch_done:
        st.info(t("step5.batch_required"))
        return

    mc_results = db.get_mc_results(inspection_id)
    if not mc_results:
        st.warning(t("step5.no_results"))
        return

    conformites = db.get_conformity_by_inspection(inspection_id)
    nc_pc_ids = [iid for iid, conf in conformites.items() if conf in ("NC", "PC")]

    if not nc_pc_ids:
        st.info(t("step5.no_nc_pc"))
        return

    df_mc = pd.DataFrame(mc_results, columns=[
        "id", "inspection_id", "item_id", "n_iterations",
        "freq_min", "freq_mode", "freq_max",
        "freq_moyenne", "freq_p05", "freq_p50", "freq_p95",
        "severite", "score_risque",
        "niveau_sev", "niveau_prob", "niveau_risque", "couleur", "created_at",
    ])
    df_mc = df_mc.sort_values("score_risque", ascending=False)
    df_nc_pc = df_mc[df_mc["item_id"].isin(nc_pc_ids)]

    if df_nc_pc.empty:
        st.info(t("step5.no_nc_pc_mc"))
        return

    st.subheader(t("step5.choose_symptom_header"))
    st.caption(t("step5.choose_symptom_caption"))

    options = []
    for _, row in df_nc_pc.iterrows():
        iid = row["item_id"]
        item_row = db.get_item(iid)
        if not item_row:
            continue
        label = t("step5.symptom_option_label",
                  id=iid, severity=row["severite"],
                  score=row["score_risque"],
                  symptom=tr_symptom(item_row)[:60])
        options.append((iid, label))

    if not options:
        st.warning(t("step5.no_valid_symptom"))
        return

    labels = [opt[1] for opt in options]
    selected_label = st.selectbox(t("step5.symptom_selector"), labels, key="pdca_symptom_selector")
    selected_idx = labels.index(selected_label)
    selected_id = options[selected_idx][0]

    item_row = db.get_item(selected_id)
    mc_row = df_nc_pc[df_nc_pc["item_id"] == selected_id].iloc[0]

    st.markdown("---")
    col1, col2 = st.columns([2, 1])
    with col1:
        st.markdown(f"**{t('step5.field_id')}** : `{selected_id}`")
        st.markdown(f"**{t('step5.field_section')}** : {item_row[1]}")
        st.markdown(f"**{t('step5.field_equipment')}** : {item_row[2]}")
        st.markdown(f"**{t('step5.field_symptom')}** : {tr_symptom(item_row)}")
    with col2:
        st.metric(t("step5.metric_severity"), item_row[4])
        st.metric(t("step5.metric_conformity"), tr_conformite(conformites.get(selected_id, "—")))
        st.metric(t("step5.metric_risk_class"), mc_row["niveau_risque"] or "—")

    st.markdown("---")

    if not ai_service.is_configured():
        st.warning(f"**{t('step4.not_configured_title')}**")
        st.info(t("step4.not_configured_body"))
        return

    st.subheader(t("step5.custom_instructions_title"))
    st.caption(t("step5.custom_instructions_caption"))

    pdca_custom = st.text_area(
        t("step5.custom_instructions_label"),
        value=st.session_state.get("pdca_custom_instructions", ""),
        height=80,
        placeholder=t("step5.custom_instructions_placeholder"),
        key="pdca_custom_input",
    )
    st.session_state.pdca_custom_instructions = pdca_custom

    st.markdown("---")

    if st.button(t("step5.generate_button", id=selected_id),
                 type="primary", key=f"gen_pdca_{selected_id}"):
        item_data = {
            "item_id": selected_id,
            "section": item_row[1],
            "equipement": item_row[2],
            "symptome": tr_symptom(item_row),
            "severity": item_row[4],
            "conformite": conformites.get(selected_id, "NC"),
        }

        with st.spinner(t("step5.generating")):
            company_profile_text = build_company_profile_text(
                st.session_state.get("company_profile_data", {})
            )
            result = ai_service.generer_plan_pdca(
                item=item_data,
                freq_moyenne=mc_row["freq_moyenne"],
                risk_class=mc_row["niveau_risque"] or "N/A",
                langue=get_current_lang(),
                company_profile=company_profile_text,
                custom_instructions=st.session_state.get("pdca_custom_instructions", ""),
            )

        if result.get("erreur"):
            st.error(t("step5.error_api", error=result["erreur"]))
        else:
            st.session_state[f"pdca_markdown_{selected_id}"] = result["data"]
            st.session_state[f"pdca_lang_{selected_id}"] = get_current_lang()
            st.success(t("step5.generated_ok"))

    pdca_key = f"pdca_markdown_{selected_id}"
    if pdca_key in st.session_state:
        st.markdown("---")
        st.subheader(t("step5.plan_generated_header"))

        from report_render import render_pdca_html
        import streamlit.components.v1 as components

        pdca_html = render_pdca_html(
            st.session_state[pdca_key],
            {
                "item_id": selected_id, "section": item_row[1], "equipement": item_row[2],
                "symptome": tr_symptom(item_row), "severity": item_row[4],
                "conformite": conformites.get(selected_id, "NC"),
            },
            {
                "freq_moyenne": mc_row["freq_moyenne"],
                "risk_class": mc_row["niveau_risque"] or "—",
                "couleur": mc_row["couleur"],
            },
            {
                "site": st.session_state.get("active_site", "—"),
                "inspection_id": inspection_id,
                "inspection_date": st.session_state.get("inspection_date_str", ""),
            },
            lang=st.session_state.get(f"pdca_lang_{selected_id}", get_current_lang()),
        )
        components.html(pdca_html, height=1000, scrolling=True)

        st.download_button(
            t("step5.download_button"),
            pdca_html.encode("utf-8"),
            f"pdca_{selected_id}.html",
            "text/html",
            key=f"download_pdca_{selected_id}",
            use_container_width=True,
        )

        st.markdown("---")
        if st.button(t("step5.regenerate_button"), key=f"regen_pdca_{selected_id}"):
            del st.session_state[pdca_key]
            st.rerun()


# ============================================================
# ÉTAPE 4 — ACTIONS CORRECTIVES (page locale à l'inspection)
# ============================================================
def render_corrective_actions():
    """Suivi des actions correctives de l'inspection courante."""
    st.markdown("---")
    st.subheader(t("actions.title"))
    st.caption(t("actions.caption"))

    conformites = db.get_conformity_by_inspection(inspection_id)
    nc_pc_ids = [i for i, c in conformites.items() if c in ("NC", "PC")]

    def _item_label(iid):
        if not iid:
            return t("actions.item_none")
        it = db.get_item(iid)
        return f"{iid} — {tr_symptom(it)[:70]}" if it else iid

    can_edit = auth.has_permission("edit_actions")
    can_delete = auth.has_permission("delete_actions")

    if not can_edit:
        render_read_only_notice()
    else:
        with st.form("corrective_action_form", clear_on_submit=True):
            item_sel = st.selectbox(t("actions.item"), [""] + nc_pc_ids, format_func=_item_label)
            action_txt = st.text_area(t("actions.action"), height=80)
            c1, c2 = st.columns(2)
            responsible = c1.text_input(t("actions.responsible"))
            resp_email = c2.text_input(t("actions.responsible_email"))
            c1, c2 = st.columns(2)
            deadline = c1.date_input(
                t("actions.deadline"), value=date.today() + timedelta(days=14), format="DD/MM/YYYY",
            )
            priority = c2.selectbox(
                t("actions.priority"), ["low", "medium", "high", "critical"], index=1,
                format_func=lambda p: t(f"actions.priority_{p}"),
            )
            submitted = st.form_submit_button(t("actions.add"), type="primary")

        if submitted:
            if not action_txt.strip():
                st.error(t("actions.err_action"))
            else:
                db.add_corrective_action(
                    inspection_id, item_sel or None, action_txt.strip(),
                    responsible.strip() or None, resp_email.strip() or None,
                    deadline.isoformat(), priority,
                )
                st.success(t("actions.saved"))

    actions = db.get_corrective_actions(inspection_id)
    if not actions:
        st.info(t("actions.empty"))
        return

    today_iso = date.today().isoformat()
    for a in actions:
        is_done = a["status"] == "done"
        late = (not is_done) and a.get("deadline") and a["deadline"] < today_iso
        badge = (f"✅ {t('actions.status_done')}" if is_done
                 else f"🔵 {t('actions.status_overdue')}" if late
                 else f"🟠 {t('actions.status_open')}")
        with st.container(border=True):
            st.markdown(f"**{badge}** — {a['action']}")
            st.caption(
                f"{t('actions.item')} : {a.get('item_id') or '—'}  |  "
                f"{t('actions.responsible')} : {a.get('responsible') or '—'}  |  "
                f"{t('actions.deadline')} : {a.get('deadline') or '—'}  |  "
                f"{t('actions.priority')} : {t('actions.priority_' + (a.get('priority') or 'medium'))}"
                + (f"  |  {t('actions.completed_on', date=a['completed_at'])}" if is_done else "")
            )
            b1, b2, _ = st.columns([1, 1, 4])
            if not is_done and can_edit and b1.button(t("actions.complete"), key=f"act_done_{a['id']}"):
                db.complete_corrective_action(a["id"])
                st.rerun()
            if can_delete and b2.button(t("actions.delete"), key=f"act_del_{a['id']}"):
                db.delete_corrective_action(a["id"])
                st.rerun()


def render_step_5_pdca():
    """Étape 4 : plan PDCA + actions correctives."""
    _render_pdca_plan()
    render_corrective_actions()


# ============================================================
# DISPATCH
# ============================================================
try:
    from ui.pages.report import render_report_page
except ImportError:
    def render_report_page(*_args, **_kwargs):
        st.header(t("nav.step_4"))
        st.info(t("step4.unavailable"))


# Garde-fou serveur : même si la session est manipulée, le rôle est revérifié.
if not auth.can_access_step(step):
    auth.deny(auth.STEP_PERMISSIONS.get(step, "view"), f"step_{step}")


# ─────────────────────────────────────────────
# Routing principal
# ─────────────────────────────────────────────
if step == STEP_COMPANY:
    render_step_company_profile()

elif step == STEP_DASHBOARD:
    # 1. Tableau de bord
    render_dashboard_page(
        db=db,
        site_name=st.session_state.get("active_site"),
    )
    st.markdown("---")
    # 2. Actions correctives juste en dessous
    render_corrective_actions_page(
        db=db,
        site_name=st.session_state.get("active_site"),
        user=current_user,
    )

elif step == STEP_COMPARISON:
    render_comparison_page(
        db=db,
        site_name=st.session_state.get("active_site"),
    )

elif step == STEP_ACTIONS_CORR:
    render_corrective_actions_page(
        db=db,
        site_name=st.session_state.get("active_site"),
        user=current_user,
    )

elif step == 0:
    render_step_0_matrix_config()
elif step == 1:
    render_step_1_inspection()
elif step == 2:
    render_step_2_qra()
elif step == 3:
    render_step_3_results()
elif step == 4:
    render_step_5_pdca()
elif step == 5:
    render_report_page(
        db,
        inspection_id,
        site_name=st.session_state.get("active_site", "—"),
    )
elif step == 6:
    _render_programme_section_picker(db.get_conformity_by_inspection(inspection_id))
    st.markdown("---")
    _prog_items = db.get_checklist_items()
    _prog_ids = None
    if db.get_domains():       # sections proposées à la planification = celles du domaine
        _prog_ids = {s_["section_id"] for s_ in _domain_sections(_inspection_domain()[0])}
    _prog_names = {}
    for it_ in _prog_items:
        _sid_ = item_section_id(it_)
        if _prog_ids is None or _sid_ in _prog_ids:
            _prog_names.setdefault(tr_section(it_), set()).add(_sid_)
    render_calendar_page(
        db,
        sites=matrice_manager.list_sites(),
        sections=sorted(_prog_names),
        default_site=st.session_state.get("active_site"),
        section_ids_by_name={n_: sorted(v_ for v_ in ids_ if v_ is not None) for n_, ids_ in _prog_names.items()},
    )
elif step == auth.STEP_ADMIN:
    render_admin_page()


# ============================================================
# FOOTER
# ============================================================
render_footer()

st.sidebar.markdown("---")
st.sidebar.caption(
    t("sidebar.footer_inspection_id", id=inspection_id)
    + "  \n"
    + t("sidebar.footer_active_matrix",
        site=st.session_state.get("active_site", "—"))
)